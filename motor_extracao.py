"""
Leitor inteligente de demonstrativos de pagamento / glosa de planos de saúde.

Uso principal (chamado pelo app):
    df, erro = extrair_glosas_de_arquivo(arquivo, plano="Bradesco Saúde", somente_glosas=True,
                                         conhecimento=motor_conhecimento.carregar())

Como funciona, em camadas (nenhuma depende do nome exato das colunas):
  1. Estrutura: acha o cabeçalho (mesmo repetido no meio da planilha), lê a data de crédito de cada
     bloco ("Liberado para crédito em 03/06/2025"), remove linhas de total/separador e junta a
     linha da guia (paciente, data, guia, protocolo) com as linhas dos procedimentos logo abaixo.
  2. Colunas: reconhece cada campo pelo nome do cabeçalho + pelo conteúdo. Quantidade nunca é
     confundida com valor. Se o arquivo tem uma única coluna de valor num extrato de pagamento,
     ela é o VALOR PAGO.
  3. Valores: o valor apresentado vem do arquivo; se o arquivo não traz, vem da tabela de valores
     cadastrada no sistema (por código TUSS ou especialidade) ou, por último, do valor que o plano
     costuma pagar por aquele código (aprendido dos próprios demonstrativos).
  4. Glosa: usa a coluna de glosa do arquivo; se não houver, deduz pela quantidade liberada menor
     que a solicitada ou pela justificativa. Diferença sem justificativa fica separada, para conferir.
  5. Duplicidade: o mesmo atendimento copiado em outra aba conta uma vez só; o mesmo atendimento
     pago em mais de um crédito é somado (complemento ou recurso) e sinalizado.

O "conhecimento" (opcional) traz o que o sistema já aprendeu: perfis de leitura confirmados pelo
usuário, tabela de valores cadastrada, valores de referência por código e significado dos códigos
de glosa (tabela TUSS 38).
"""
import csv
import hashlib
import io
import itertools
import re
import unicodedata
import xml.etree.ElementTree as ET
from difflib import SequenceMatcher

import pandas as pd

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import docx
except ImportError:
    docx = None


# Colunas originais (compatibilidade com o restante do sistema)
COLUNAS_PADRAO_GLOSAS = [
    "Data do evento",
    "Protocolo",
    "Doc Original",
    "Nome do paciente",
    "Código",
    "Procedimento",
    "Valor informado",
    "Valor pago",
    "Valor da glosa",
    "Motivo da glosa",
]
# Colunas novas (ficam sempre depois das originais)
COLUNAS_EXTRAS = ["Quantidade", "Quantidade liberada", "Valor unitário", "Lote", "Status da glosa",
                  "Data de pagamento", "Código da glosa", "Valor de referência", "Fonte do valor informado",
                  "Diferença sem justificativa", "Alertas"]
COLUNAS_COMPLETAS = COLUNAS_PADRAO_GLOSAS + COLUNAS_EXTRAS

PLANO_AUTOMATICO = "Identificar pelo arquivo"
TOL = 0.02

FONTE_ARQUIVO = "Demonstrativo"
FONTE_CADASTRO = "Tabela cadastrada"
FONTE_APRENDIDO = "Valor padrão pago pelo plano"
FONTE_NENHUMA = "Não identificado"

# --------------------------------------------------------------------------
# Configuração dos campos
# --------------------------------------------------------------------------
CAMPOS = {
    "Data do evento": dict(
        tipo="data",
        sin=["DATA ATENDIMENTO", "DATA REALIZACAO", "DATA EXECUCAO", "DATA EVENTO",
             "DATA PROCEDIMENTO", "DATA SERVICO", "DT ATENDIMENTO", "DT REALIZACAO",
             "DATA", "DT"],
        excl=["NASC", "PAGAMENTO", "PGTO", "CREDITO", "VENC", "PROTOCOLO", "ENVIO", "EMISSAO", "ALTA"],
    ),
    "Protocolo": dict(
        tipo="id",
        sin=["NUMERO PROTOCOLO", "NRO PROTOCOLO", "NR PROTOCOLO", "N PROTOCOLO",
             "NUM PROTOCOLO", "PROTOCOLO", "PROTOC"],
        excl=["DATA", "DT", "VALOR", "STATUS", "SITUACAO"],
    ),
    "Doc Original": dict(
        tipo="id",
        sin=["DOC ORIGINAL", "DOCUMENTO ORIGINAL", "NUMERO GUIA PRESTADOR", "GUIA PRESTADOR",
             "NUMERO GUIA", "NRO GUIA", "NR GUIA", "N GUIA", "NUM GUIA", "GUIA",
             "NUMERO DOCUMENTO", "NRO DOCUMENTO", "DOCUMENTO", "DOC", "SENHA"],
        excl=["DATA", "DT", "VALOR", "TIPO", "CPF", "CNPJ", "CARTEIRA"],
    ),
    "Nome do paciente": dict(
        tipo="nome",
        sin=["NOME PACIENTE", "NOME BENEFICIARIO", "PACIENTE", "BENEFICIARIO",
             "SEGURADO", "USUARIO", "CLIENTE", "NOME"],
        excl=["MEDICO", "PROFISSIONAL", "PLANO", "OPERADORA", "PRESTADOR", "EXECUTANTE",
              "SOLICITANTE", "CONVENIO", "MAE", "RESPONSAVEL", "PROCEDIMENTO", "HOSPITAL",
              "CLINICA", "SERVICO"],
    ),
    "Código": dict(
        tipo="codigo",
        sin=["CODIGO PROCEDIMENTO", "COD PROCEDIMENTO", "COD PROC", "CODIGO TUSS",
             "COD TUSS", "TUSS", "CODIGO", "COD", "AMB", "ITEM", "PROC"],
        excl=["GLOSA", "MOTIVO", "GUIA", "BENEFICIARIO", "CARTEIRA", "PACIENTE", "SENHA",
              "PROTOCOLO", "LOTE", "MEDICO", "CRM", "CPF", "CNPJ", "TABELA"],
    ),
    "Procedimento": dict(
        tipo="texto",
        sin=["DESCRICAO PROCEDIMENTO", "NOME PROCEDIMENTO", "PROCEDIMENTO", "DESCRICAO",
             "SERVICO", "EVENTO", "TRATAMENTO", "EXAME"],
        excl=["GLOSA", "MOTIVO", "CODIGO", "COD", "DATA", "VALOR"],
    ),
    "Valor informado": dict(
        tipo="dinheiro",
        sin=["VALOR INFORMADO", "VALOR APRESENTADO", "VALOR COBRADO", "VALOR FATURADO",
             "VALOR PLANO", "INFORMADO", "APRESENTADO", "COBRADO", "FATURADO",
             "VALOR", "VLR", "VL"],
        excl=["GLOSA", "PAGO", "RECEBIDO", "LIBERADO", "LIQUIDO", "RETIDO", "APROVADO",
              "PAGAMENTO", "UNIT", "UNITARIO", "QTD", "QTDE", "QUANTIDADE", "SOLICITADA", "SESSAO",
              "PROCESSADO"],
    ),
    "Valor pago": dict(
        tipo="dinheiro",
        sin=["VALOR PAGO", "VALOR LIBERADO", "VALOR RECEBIDO", "VALOR LIQUIDO",
             "VALOR APROVADO", "VALOR ACEITO", "VL PAGO", "VLR PAGO", "PAGO",
             "RECEBIDO", "LIBERADO", "LIQUIDO", "APROVADO"],
        excl=["GLOSA", "INFORMADO", "APRESENTADO", "COBRADO", "QTD", "QTDE", "QUANTIDADE",
              "SOLICITADA", "SESSAO", "DATA", "DT", "UNIT", "UNITARIO", "PARA"],
    ),
    "Valor da glosa": dict(
        tipo="dinheiro",
        sin=["VALOR GLOSA", "VALOR GLOSADO", "VL GLOSA", "VLR GLOSA", "VALOR RETIDO",
             "GLOSADO", "GLOSA", "RETIDO"],
        excl=["MOTIVO", "JUSTIFICATIVA", "DESCRICAO", "CODIGO", "COD", "OBS", "TIPO", "QTD",
              "QTDE", "QUANTIDADE", "STATUS", "SITUACAO", "RELACAO"],
    ),
    "Motivo da glosa": dict(
        tipo="texto",
        sin=["MOTIVO GLOSA", "DESCRICAO GLOSA", "DESCRICAO MOTIVO", "MOTIVO",
             "JUSTIFICATIVA", "OBSERVACAO", "OBS", "CRITICA", "MENSAGEM"],
        excl=["VALOR", "VL", "VLR", "STATUS", "SITUACAO", "TEM", "HOUVE", "CODIGO", "COD", "TIPO"],
    ),
    "Quantidade": dict(
        tipo="qtd",
        sin=["QUANTIDADE SOLICITADA", "QTD SOLICITADA", "QTDE SOLICITADA", "QUANT SOLICITADA",
             "QUANTIDADE COBRADA", "QTD COBRADA", "QTD APRESENTADA", "QUANTIDADE APRESENTADA",
             "QUANTIDADE EXECUTADA", "QTD EXECUTADA", "QTDE EXECUTADA", "QTD REALIZADA",
             "QUANTIDADE", "QTDE", "QTD", "SESSOES", "QUANT"],
        excl=["VALOR", "VL", "VLR", "GLOSA", "GLOSADA", "LIBERAD", "PAGA", "APROVAD", "AUTORIZAD"],
    ),
    "Quantidade liberada": dict(
        tipo="qtd",
        sin=["QUANTIDADE LIBERADA", "QTD LIBERADA", "QTDE LIBERADA", "QUANT LIBERADA",
             "QUANTIDADE PAGA", "QTD PAGA", "QTDE PAGA", "QUANTIDADE APROVADA", "QTD APROVADA",
             "QUANTIDADE AUTORIZADA", "QTD AUTORIZADA", "QTD LIBERADO", "QUANTIDADE LIBERADO"],
        excl=["VALOR", "VL", "VLR"],
        exige_nome=True,
    ),
    "Valor unitário": dict(
        tipo="dinheiro",
        sin=["VALOR UNITARIO", "VLR UNITARIO", "VL UNITARIO", "VALOR UNIT", "UNITARIO", "UNIT"],
        excl=["GLOSA", "PAGO", "TOTAL", "QTD", "QTDE", "QUANTIDADE"],
    ),
    "Lote": dict(
        tipo="id",
        sin=["NUMERO LOTE PRESTADOR", "NUMERO LOTE", "NRO LOTE", "NR LOTE", "N LOTE", "LOTE"],
        excl=["DATA", "DT", "VALOR", "STATUS", "SITUACAO"],
    ),
    "Status da glosa": dict(
        tipo="texto",
        sin=["STATUS GLOSA", "SITUACAO GLOSA", "TEM GLOSA", "HOUVE GLOSA", "GLOSADO SN",
             "SITUACAO", "STATUS"],
        excl=["VALOR", "VL", "VLR", "DATA", "DT", "MOTIVO", "DESCRICAO", "CODIGO", "COD"],
    ),
    "Data de pagamento": dict(
        tipo="data",
        sin=["DATA PAGAMENTO", "DATA DO PAGAMENTO", "DT PAGAMENTO", "DATA PGTO", "DT PGTO",
             "DATA CREDITO", "DATA LIQUIDACAO", "DATA QUITACAO", "PAGAMENTO"],
        excl=["ATENDIMENTO", "REALIZACAO", "EXECUCAO", "EVENTO", "NASC", "PROTOCOLO", "ENVIO",
              "EMISSAO", "ALTA", "VENC", "VALOR", "FORMA"],
        exige_nome=True,
    ),
    "Código da glosa": dict(
        tipo="livre",
        sin=["CODIGO GLOSA", "COD GLOSA", "TIPO GLOSA", "CODIGO MOTIVO GLOSA", "CODIGO MOTIVO",
             "COD MOTIVO", "CODIGO DA GLOSA", "CODIGO MENSAGEM"],
        excl=["VALOR", "VL", "VLR", "DESCRICAO", "PROCEDIMENTO"],
        exige_nome=True,
    ),
}

CAMPOS_DINHEIRO = ["Valor informado", "Valor pago", "Valor da glosa"]
CAMPOS_QTD = ["Quantidade", "Quantidade liberada"]
LIMIAR_ATRIBUICAO = 35
MAX_COLUNAS_NUMERICAS_LIVRES = 8  # limita as permutações testadas (evita lentidão)

# --------------------------------------------------------------------------
# Utilidades de texto
# --------------------------------------------------------------------------
_RE_DATA = re.compile(r"\b(\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}|\d{4}-\d{2}-\d{2})\b")
_RE_DIN = re.compile(r"^(R\$\s*)?-?\(?\d{1,3}(\.\d{3})*,\d{1,2}\)?$|^(R\$\s*)?-?\(?\d+[.,]\d{1,2}\)?$")
_RE_COD = re.compile(r"^(\d{6,10}|\d{1,2}\.\d{2}\.\d{2}\.\d{2}-?\d?|\d{2}\.\d{3}\.\d{3}-?\d?)$")
_RE_NOME = re.compile(r"^[^\W\d_]+(?:[ '.-]+[^\W\d_]+)+$")
_RE_DIN_TEXTO = re.compile(r"\d{1,3}(?:\.\d{3})*,\d{2}")
# Protocolo/guia no texto corrido: exige pelo menos um dígito para não capturar palavras
_RE_PROTOCOLO = re.compile(r"(?:protocolo|prot\.)[\s:nº°.]*((?=[A-Z0-9.-]*\d)[A-Z0-9.-]{5,20})", re.I)
_RE_GUIA = re.compile(r"(?:guia|senha|doc\.)[\s:nº°.]*((?=[A-Z0-9.-]*\d)[A-Z0-9.-]{5,20})", re.I)
_RE_LOTE = re.compile(r"lote[\s:nº°.]*((?=[A-Z0-9.-]*\d)[A-Z0-9.-]{3,20})", re.I)
_RE_QTD = re.compile(r"^\d{1,3}([.,]0+)?$")
_RE_CTX_PAGAMENTO = re.compile(r"(PAGAMENTO|PGTO|CREDITO|LIBERADO|QUITACAO|DEPOSITO)")
_RE_COD_GLOSA = re.compile(r"^\s*(\d{4})\s*[-–:.]?\s*(.*)$")
_PALAVRAS_PROC = {"CONSULTA", "EXAME", "SESSAO", "PROCEDIMENTO", "TERAPIA", "HEMOGRAMA",
                  "RADIOGRAFIA", "ULTRASSONOGRAFIA", "TOMOGRAFIA", "RESSONANCIA", "CIRURGIA",
                  "AVALIACAO", "ATENDIMENTO", "DIARIA", "TAXA", "MATERIAL", "MEDICAMENTO"}


def _norm(s):
    s = unicodedata.normalize("NFKD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c)).upper()
    s = re.sub(r"[^A-Z0-9]+", " ", s)
    return s.strip()


def _txt(v):
    if v is None:
        return ""
    try:
        if pd.isnull(v):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(v, float) and v.is_integer() and abs(v) >= 100000:
        return str(int(v))
    s = str(v).strip()
    return "" if s.lower() in ("nan", "none", "nat") else s


def _id_txt(v):
    """Identificadores (guia, protocolo, código): sem '.0' de números lidos do Excel."""
    s = _txt(v)
    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]
    return s


def limpar_valor_monetario(val):
    try:
        if pd.isnull(val):
            return 0.0
    except (TypeError, ValueError):
        pass
    if isinstance(val, (int, float)):
        return float(val)
    if hasattr(val, "year") and hasattr(val, "month"):  # Excel transformou o número em data
        return 0.0
    s = str(val).strip()
    negativo = (s.startswith("(") and s.endswith(")")) or s.startswith("-")
    s = s.replace("R$", "").strip()
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    s = re.sub(r"[^0-9.]", "", s)
    try:
        v = float(s) if s else 0.0
    except ValueError:
        v = 0.0
    return -v if negativo else v


def _numero_ou_nan(val):
    s = _txt(val)
    if not s:
        return float("nan")
    return limpar_valor_monetario(val)


_RE_DATA_BR = re.compile(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})$")
_RE_DATA_ISO = re.compile(r"^(\d{4})-(\d{2})-(\d{2})")


def padronizar_data(val):
    if hasattr(val, "strftime") and not isinstance(val, str):
        try:
            return "-" if pd.isna(val) else val.strftime("%d/%m/%Y")
        except (ValueError, TypeError):
            return "-"
    s = _txt(val)
    if not s:
        return "-"
    m = _RE_DATA_ISO.match(s)
    if m and 1 <= int(m.group(2)) <= 12 and 1 <= int(m.group(3)) <= 31:
        return f"{m.group(3)}/{m.group(2)}/{m.group(1)}"
    m = _RE_DATA_BR.match(s)
    if m and 1 <= int(m.group(2)) <= 12 and 1 <= int(m.group(1)) <= 31:
        return f"{int(m.group(1)):02d}/{int(m.group(2)):02d}/{m.group(3)}"
    try:
        if re.match(r"^\d{4}-\d{2}-\d{2}", s):
            dt = pd.to_datetime(s[:10], errors="coerce")
        else:
            m = _RE_DATA.search(s)
            dt = pd.to_datetime(m.group(1) if m else s, errors="coerce", dayfirst=True)
        if pd.notnull(dt):
            return dt.strftime("%d/%m/%Y")
    except Exception:
        pass
    return "-"


_CACHE_SCORE = {}


def _score_nome(h_norm, campo, campos=None):
    if not h_norm:
        return 0.0
    chave = (id(campos) if campos is not None else 0, h_norm, campo)
    if chave in _CACHE_SCORE:
        return _CACHE_SCORE[chave]
    if len(_CACHE_SCORE) > 200_000:
        _CACHE_SCORE.clear()
    r = _CACHE_SCORE[chave] = _score_nome_calc(h_norm, campo, campos)
    return r


def _score_nome_calc(h_norm, campo, campos=None):
    cfg = (campos or CAMPOS)[campo]
    hw = h_norm.split()
    hc = h_norm.replace(" ", "")
    if any(w.startswith(e) for w in hw for e in cfg["excl"]):
        return 0.0
    melhor = 0.0
    for termo in cfg["sin"]:
        tw = termo.split()
        tc = termo.replace(" ", "")
        if h_norm == termo:
            s = 100.0
        elif hc == tc:
            s = 95.0
        elif all(w in hw for w in tw):
            s = min(95.0, 70.0 + 8 * len(tw))
        elif len(tc) >= 6 and tc in hc:
            s = 75.0
        elif len(tc) >= 5 and 2 * min(len(hc), len(tc)) / (len(hc) + len(tc)) >= 0.84:
            r = SequenceMatcher(None, hc, tc).ratio()
            s = 60.0 * r if r >= 0.84 else 0.0
        else:
            s = 0.0
        melhor = max(melhor, s)
    return melhor


def _eh_qtd(orig, v):
    if isinstance(orig, (int, float)) and not isinstance(orig, bool):
        try:
            return float(orig).is_integer() and 0 <= abs(float(orig)) <= 999
        except (TypeError, ValueError):
            return False
    return bool(_RE_QTD.match(v.replace(" ", "")))


def _perfil(serie):
    pares = [(v, _txt(v)) for v in serie.tolist()]
    pares = [(v, t) for v, t in pares if t][:60]
    amostra = [t for _, t in pares]
    perfil = dict(data=0.0, dinheiro=0.0, codigo=0.0, nome=0.0, texto=0.0, id=0.0, qtd=0.0,
                  tam=0.0, n=len(amostra), max=0.0)
    if not amostra:
        return perfil
    n = len(amostra)
    c = dict(data=0, dinheiro=0, codigo=0, nome=0, texto=0)
    for orig, v in pares:
        if isinstance(orig, (int, float)) and not isinstance(orig, bool):
            if float(orig).is_integer() and abs(orig) >= 100000:
                c["codigo"] += 1
            else:
                c["dinheiro"] += 1
        elif _RE_DATA.search(v) and not _RE_DIN.match(v):
            c["data"] += 1
        elif _RE_COD.match(v.replace(" ", "")):
            c["codigo"] += 1
        elif _RE_DIN.match(v):
            c["dinheiro"] += 1
        elif re.fullmatch(r"-?\d{1,5}", v):
            c["dinheiro"] += 1
        elif any(ch.isalpha() for ch in v):
            c["texto"] += 1
            palavras = _norm(v).split()
            if (_RE_NOME.match(v) and 6 <= len(v) <= 70 and 2 <= len(palavras) <= 6
                    and not (set(palavras) & _PALAVRAS_PROC)):
                c["nome"] += 1
    for k, v in c.items():
        perfil[k] = v / n
    perfil["qtd"] = sum(1 for orig, v in pares if _eh_qtd(orig, v)) / n
    perfil["tam"] = sum(len(v) for v in amostra) / n
    perfil["id"] = sum(
        1 for v in amostra
        if re.fullmatch(r"[\w./-]{4,25}", v) and any(ch.isdigit() for ch in v)
        and not _RE_DATA.search(v) and not _RE_DIN.match(v)
    ) / n
    nums = [abs(limpar_valor_monetario(o)) for o, v in pares if re.fullmatch(r"[\d.,R$ ()-]+", v)]
    perfil["max"] = max(nums) if nums else 0.0
    return perfil


def _score_campo(campo, h_norm, perfil, campos=None):
    campos = campos or CAMPOS
    cfg = campos[campo]
    tipo = cfg["tipo"]
    frac = 1.0 if tipo == "livre" else perfil.get(tipo, 0.0)  # "livre": decide só pelo nome do cabeçalho
    nome = _score_nome(h_norm, campo, campos)
    if cfg.get("exige_nome") and nome <= 0:
        return 0.0
    compat = 0.3 + 0.7 * frac
    score = nome * compat
    if tipo in ("data", "codigo", "nome"):
        score += 50.0 * frac
    elif tipo == "texto":
        score += 25.0 * frac
    return score


# --------------------------------------------------------------------------
# Estrutura da tabela
# --------------------------------------------------------------------------
def _linha_eh_ruido(valores):
    """Separadores ('-----'), rodapés e linhas sem conteúdo útil."""
    textos = [t for t in (_txt(v) for v in valores) if t]
    if not textos:
        return True
    return any(re.fullmatch(r"[-=_*.]{5,}.*", t) for t in textos)


def _preparar_tabela(raw, campos=None, com_origem=False):
    """Acha o cabeçalho e devolve a tabela com os nomes das colunas.
    com_origem=True acrescenta '__orig' (número da linha na planilha original)."""
    campos = campos or CAMPOS
    raw = raw.copy().dropna(how="all").dropna(axis=1, how="all")
    origem = list(raw.index)
    raw = raw.reset_index(drop=True)
    raw.columns = range(raw.shape[1])
    if raw.empty:
        return raw

    melhor_i, melhor_hits = None, 0
    for i in range(min(len(raw), 30)):
        hits = sum(
            1 for cel in raw.iloc[i].tolist()
            if _txt(cel) and not re.fullmatch(r"[\d.,/ R$-]+", _txt(cel))
            and any(_score_nome(_norm(_txt(cel)), campo, campos) >= 70 for campo in campos)
        )
        if hits > melhor_hits:
            melhor_i, melhor_hits = i, hits

    if melhor_hits >= 2:
        cab_i = melhor_i
    else:
        primeira = [_txt(c) for c in raw.iloc[0].tolist() if _txt(c)]
        if primeira and all(not re.fullmatch(r"[\d.,/ R$-]+", t) for t in primeira) and len(primeira) >= 3:
            cab_i = 0
        else:
            cab_i = None

    if cab_i is None:
        df = raw.copy()
        df.columns = [f"COL_{i}" for i in range(df.shape[1])]
        if com_origem:
            df["__orig"] = origem
        return _remover_totais(df)

    cabecalho = [_txt(c) for c in raw.iloc[cab_i].tolist()]
    nomes, usados = [], {}
    for j, cel in enumerate(cabecalho):
        n = cel or f"COL_{j}"
        usados[n] = usados.get(n, 0) + 1
        nomes.append(n if usados[n] == 1 else f"{n}_{usados[n]}")
    df = raw.iloc[cab_i + 1:].copy()
    if com_origem:
        df["__orig"] = origem[cab_i + 1:]
    df = df.reset_index(drop=True)
    df.columns = nomes + (["__orig"] if com_origem else [])

    # cabeçalho repetido no meio (planilhas com vários blocos) e linhas separadoras
    cab_norm = {_norm(c) for c in cabecalho if c}
    textos = df[nomes].astype(object).where(df[nomes].notna(), "").astype(str).apply(lambda c: c.str.strip())
    unicos = pd.unique(textos.values.ravel())
    eh_cab = {u: (bool(u) and _norm(u) in cab_norm) for u in unicos}
    repetido = textos.apply(lambda c: c.map(eh_cab)).sum(axis=1) >= 2
    separador = textos.apply(lambda c: c.str.match(r"^[-=_*.]{5,}")).any(axis=1)
    vazio = textos.isin(["", "nan", "None", "NaT"]).all(axis=1)
    df = df[~(repetido | separador | vazio)].reset_index(drop=True)
    return _remover_totais(df)


def _remover_totais(df):
    def eh_total(linha):
        for cel in linha.tolist():
            if re.match(r"^(SUB )?TOTAL( GERAL)?$|^TOTAIS$|^TOTAL (GERAL|DA|DO|DOS|DAS)\b", _norm(_txt(cel))):
                return True
        return False

    if df.empty:
        return df
    return df[~df.apply(eh_total, axis=1)].reset_index(drop=True)


def _texto_global(raw, rotulo, regex_valor, max_linhas=30):
    """Procura, no topo da planilha (fora da tabela), valores soltos como
    'Data de pagamento: 10/09/2026' ou 'Lote: 12345'. Retorna o texto encontrado ou ''."""
    try:
        n = min(len(raw), max_linhas)
        for i in range(n):
            cels = [_txt(c) for c in raw.iloc[i].tolist()]
            for j, t in enumerate(cels):
                if not t or not re.search(rotulo, _norm(t)):
                    continue
                m = regex_valor.search(t)
                if m:
                    return m.group(1) if m.groups() else m.group(0)
                for prox in [c for c in cels[j + 1:j + 4] if c]:
                    m = regex_valor.search(prox)
                    if m:
                        return m.group(1) if m.groups() else m.group(0)
    except Exception:
        pass
    return ""


_RE_LOTE_CEL = re.compile(r"^((?=[A-Z0-9.-]*\d)[A-Z0-9.-]{3,20})$", re.I)


def _lote_global(raw):
    achado = _texto_global(raw, r"\bLOTE\b", _RE_LOTE)
    if not achado:
        achado = _texto_global(raw, r"\bLOTE\b", _RE_LOTE_CEL)
    return achado


def _data_pagamento_global(raw):
    achado = _texto_global(raw, r"\b(PAGAMENTO|PGTO|CREDITO)\b", _RE_DATA)
    return padronizar_data(achado) if achado else ""


def _pagamento_por_linha(raw):
    """Data de pagamento/crédito que vale para cada linha da planilha original: planilhas em blocos
    trazem 'Liberado para crédito em 03/06/2025' antes de cada bloco."""
    atual, res = "", {}
    for idx, linha in raw.iterrows():
        textos = [_txt(v) for v in linha.tolist()]
        juntos = " ".join(t for t in textos if t)
        if juntos and _RE_CTX_PAGAMENTO.search(_norm(juntos)):
            m = _RE_DATA.search(juntos)
            if m and len([t for t in textos if t and not re.fullmatch(r"[-=_*.]{5,}.*", t)]) <= 4:
                atual = padronizar_data(m.group(1))
        res[idx] = atual
    return res


def _tipo_documento(raw, max_linhas=12):
    topo = " ".join(_norm(_txt(v)) for i in range(min(len(raw), max_linhas)) for v in raw.iloc[i].tolist())
    if re.search(r"\b(EXTRATO|DEMONSTRATIVO) DE PAGAMENTO\b|\bLIBERADO PARA CREDITO\b|\bEXTRATO\b", topo):
        return "pagamento"
    if re.search(r"\bANALISE DE CONTA\b|\bGLOSA", topo):
        return "analise"
    return ""


def assinatura_colunas(colunas):
    """Identifica um layout de arquivo pelo conjunto de nomes de coluna (para lembrar mapeamentos)."""
    nomes = sorted({_norm(c) for c in colunas if c and not str(c).startswith(("COL_", "__"))})
    return hashlib.sha1("|".join(nomes).encode()).hexdigest()[:16] if nomes else ""


# --------------------------------------------------------------------------
# Mapeamento das colunas
# --------------------------------------------------------------------------
def _serie_num(df, col, cache):
    """Converte a coluna para número uma única vez."""
    if col not in cache:
        cache[col] = df[col].apply(limpar_valor_monetario)
    return cache[col]


def _consistencia(df, c_inf, c_pago, c_glosa, cache):
    if len(df) == 0:
        return 0.0
    a = _serie_num(df, c_inf, cache)
    b = _serie_num(df, c_pago, cache)
    g = _serie_num(df, c_glosa, cache)
    return float(((a - b - g).abs() < 0.02).mean())


def _descobrir_mapeamento(df):
    cols = [c for c in df.columns if not str(c).startswith("__")]
    perfis = {c: _perfil(df[c]) for c in cols}
    hnorm = {c: _norm(c) for c in cols}

    candidatos = []
    for campo in CAMPOS:
        for c in cols:
            s = _score_campo(campo, hnorm[c], perfis[c])
            if s >= LIMIAR_ATRIBUICAO:
                outro = max((_score_nome(hnorm[c], o) for o in CAMPOS if o != campo), default=0.0)
                if outro >= 70 and _score_nome(hnorm[c], campo) < LIMIAR_ATRIBUICAO:
                    continue
                # quantidade nunca vira valor em dinheiro (e vice-versa)
                if campo in CAMPOS_DINHEIRO + ["Valor unitário"] and max(
                        _score_nome(hnorm[c], q) for q in CAMPOS_QTD) >= 70:
                    continue
                candidatos.append((s, campo, c))
    candidatos.sort(key=lambda x: -x[0])

    mapa, usadas = {}, set()
    for s, campo, c in candidatos:
        if campo not in mapa and c not in usadas:
            mapa[campo] = c
            usadas.add(c)

    # Colunas de dinheiro sem nome reconhecível: testa combinações que fecham a conta
    # apresentado = pago + glosa. Só entram colunas cujo nome não indica outro campo
    # e que não parecem contagem (0, 1, 2...).
    faltam = [f for f in CAMPOS_DINHEIRO if f not in mapa]
    livres = [
        c for c in cols
        if c not in usadas and perfis[c]["dinheiro"] >= 0.6
        and not (perfis[c]["qtd"] >= 0.9 and perfis[c]["max"] <= 50)
        and max((_score_nome(hnorm[c], o) for o in CAMPOS if o not in CAMPOS_DINHEIRO), default=0.0) < 70
    ][:MAX_COLUNAS_NUMERICAS_LIVRES]
    if faltam and livres and (len(mapa.keys() & set(CAMPOS_DINHEIRO)) + len(livres) >= 2):
        cache = {}
        melhor, melhor_score = None, -1.0
        for perm in itertools.permutations(livres, min(len(faltam), len(livres))):
            tentativa = dict(zip(faltam, perm))
            final = {**mapa, **tentativa}
            score = 0.0
            if all(f in final for f in CAMPOS_DINHEIRO):
                score += 10 * _consistencia(
                    df, final["Valor informado"], final["Valor pago"], final["Valor da glosa"], cache
                )
            pos = [cols.index(final[f]) for f in CAMPOS_DINHEIRO if f in final]
            score += 0.1 * sum(1 for a, b in zip(pos, pos[1:]) if a < b)
            if score > melhor_score:
                melhor, melhor_score = tentativa, score
        if melhor:
            for f, c in melhor.items():
                mapa[f] = c
                usadas.add(c)

    # tabela sem cabeçalho (COL_0, COL_1...): usa o conteúdo para achar textos
    sobras = [c for c in cols if c not in usadas and str(c).startswith("COL_") and perfis[c]["texto"] >= 0.6]
    for campo in ("Procedimento", "Motivo da glosa"):
        if campo not in mapa and sobras:
            c = sobras.pop(0)
            mapa[campo] = c
            usadas.add(c)

    return mapa


def _status_indica_glosa(valor):
    """True = o status diz que houve glosa; False = diz que não houve; None = não dá para saber."""
    n = _norm(valor)
    if not n:
        return None
    if re.search(r"\bSEM GLOSA\b|\bNAO GLOSAD|\bNAO HOUVE\b|\bSEM RESTRICAO\b", n):
        return False
    if n in ("NAO", "N"):
        return False
    if n in ("SIM", "S"):
        return True
    if re.search(r"GLOS|NEGAD|RECUS|INDEFERID|PARCIAL|CORTAD|DEVOLVID", n):
        return True
    if re.search(r"\bPAGO\b|LIBERAD|APROVAD|ACEIT|AUTORIZAD|QUITAD|\bOK\b", n):
        return False
    return None


# --------------------------------------------------------------------------
# Montagem das linhas (antes da análise de glosa)
# --------------------------------------------------------------------------
_MARCADORES = ["_informado_no_arquivo", "_glosa_no_arquivo", "_status_glosa", "_tabela", "_linha_guia"]


def _montar_saida(df, mapa, pag_global="", lote_global="", pag_por_linha=None):
    def col(campo):
        return df[mapa[campo]] if campo in mapa else None

    out = pd.DataFrame(index=df.index)
    s = col("Data do evento")
    out["Data do evento"] = s.apply(padronizar_data) if s is not None else pd.NA
    for campo in ("Protocolo", "Doc Original", "Lote", "Código"):
        s = col(campo)
        out[campo] = s.apply(_id_txt) if s is not None else ("" if campo == "Código" else pd.NA)
    s = col("Nome do paciente")
    out["Nome do paciente"] = s.apply(_txt) if s is not None else pd.NA
    if col("Lote") is None and lote_global:
        out["Lote"] = lote_global

    s = col("Data de pagamento")
    if s is not None:
        out["Data de pagamento"] = s.apply(padronizar_data)
    elif pag_por_linha is not None and "__orig" in df.columns and df["__orig"].map(pag_por_linha).fillna("").ne("").any():
        out["Data de pagamento"] = df["__orig"].map(pag_por_linha).fillna("")
    else:
        out["Data de pagamento"] = pag_global if pag_global else pd.NA

    s = col("Procedimento")
    out["Procedimento"] = s.apply(_txt) if s is not None else ""
    out["Motivo da glosa"] = col("Motivo da glosa").apply(_txt) if "Motivo da glosa" in mapa else ""
    out["Código da glosa"] = col("Código da glosa").apply(_id_txt) if "Código da glosa" in mapa else ""

    out["Valor informado"] = col("Valor informado").apply(_numero_ou_nan) if "Valor informado" in mapa else float("nan")
    out["Valor pago"] = col("Valor pago").apply(limpar_valor_monetario) if "Valor pago" in mapa else float("nan")
    out["Valor da glosa"] = col("Valor da glosa").apply(_numero_ou_nan) if "Valor da glosa" in mapa else float("nan")
    out["Quantidade"] = col("Quantidade").apply(_numero_ou_nan) if "Quantidade" in mapa else float("nan")
    out["Quantidade liberada"] = (col("Quantidade liberada").apply(_numero_ou_nan)
                                  if "Quantidade liberada" in mapa else float("nan"))
    out["Valor unitário"] = (col("Valor unitário").apply(limpar_valor_monetario).abs()
                             if "Valor unitário" in mapa else float("nan"))
    s = col("Status da glosa")
    out["_status_glosa"] = s.apply(_status_indica_glosa) if s is not None else None

    # Linha da guia x linhas dos procedimentos: quando a tabela tem código de procedimento em boa
    # parte das linhas, as linhas SEM código e SEM descrição que são seguidas de procedimentos são
    # o "cabeçalho da guia" (paciente, data, guia, protocolo e total). Elas passam os dados para os
    # procedimentos de baixo e não contam como atendimento.
    tem_cod = out["Código"].fillna("").astype(str).str.strip().ne("") | out["Procedimento"].fillna("").astype(str).str.strip().ne("")
    out["_linha_guia"] = False
    if ("Código" in mapa or "Procedimento" in mapa) and 0.2 <= tem_cod.mean() < 1.0:
        prox_tem_cod = tem_cod.shift(-1, fill_value=False)
        out["_linha_guia"] = (~tem_cod) & prox_tem_cod

    # Preenchimento para baixo (dados hierárquicos)
    for c in ["Data do evento", "Protocolo", "Doc Original", "Nome do paciente", "Lote", "Data de pagamento"]:
        out[c] = out[c].replace(["", "-", "Não identificado"], pd.NA).ffill()
        out[c] = out[c].fillna("Não identificado" if c == "Nome do paciente" else "-")

    out = out[~out["_linha_guia"]].copy()
    out["_informado_no_arquivo"] = "Valor informado" in mapa
    out["_glosa_no_arquivo"] = "Valor da glosa" in mapa
    return out


def _limpar_linhas(out):
    if out.empty:
        return out
    num = lambda c: pd.to_numeric(out[c], errors="coerce").fillna(0)
    vazio = (
        (out["Nome do paciente"].isin(["", "Não identificado"])) & (out["Código"].fillna("").isin(["", "-"]))
        & (num("Valor informado") == 0) & (num("Valor pago") == 0) & (num("Valor da glosa") == 0)
    )
    ref = (out["Nome do paciente"].fillna("") + " " + out["Procedimento"].fillna("")).map(_norm)
    total = ref.str.match(r"^(SUB ?)?TOTAL\b|^TOTAIS\b")
    return out[~(vazio | total)].reset_index(drop=True)


def _processar_bruto(raw, perfis=None, nome_tabela=""):
    """Lê uma tabela bruta (aba, tabela de PDF...). Retorna (linhas, mapa, info)."""
    vazio = pd.DataFrame()
    info = dict(tabela=nome_tabela, colunas=[], assinatura="", mapa={}, origem_mapa="automático",
                tipo_documento="", linhas=0, linhas_guia=0, observacoes=[])
    if raw is None or raw.empty:
        return vazio, {}, info
    pag_global = _data_pagamento_global(raw)
    lote_global = _lote_global(raw)
    tipo_doc = _tipo_documento(raw)
    pag_por_linha = _pagamento_por_linha(raw)
    tabela = _preparar_tabela(raw, com_origem=True)
    if tabela.empty or tabela.shape[1] < 3:
        return vazio, {}, info
    colunas = [c for c in tabela.columns if not str(c).startswith("__")]
    assinatura = assinatura_colunas(colunas)
    info.update(colunas=colunas, assinatura=assinatura, tipo_documento=tipo_doc)

    perfil_salvo = (perfis or {}).get(assinatura)
    if perfil_salvo:
        mapa = {campo: c for campo, c in perfil_salvo.items() if c in colunas and campo in CAMPOS}
        info["origem_mapa"] = "perfil salvo"
    else:
        mapa = _descobrir_mapeamento(tabela)
        # Extrato de pagamento com uma única coluna de valor: ela é o VALOR PAGO
        dinheiro = [f for f in CAMPOS_DINHEIRO if f in mapa]
        if dinheiro == ["Valor informado"] and (
                tipo_doc == "pagamento" or "Quantidade liberada" in mapa
                or _norm(mapa["Valor informado"]) in ("VALOR", "VALOR R", "VLR", "VL", "VALOR RS")):
            mapa["Valor pago"] = mapa.pop("Valor informado")
            info["observacoes"].append("A única coluna de valor do arquivo foi lida como VALOR PAGO.")
    info["mapa"] = dict(mapa)
    if len(mapa) < 2 or not (set(mapa) & set(CAMPOS_DINHEIRO)):
        return vazio, mapa, info

    out = _montar_saida(tabela, mapa, pag_global, lote_global, pag_por_linha)
    info["linhas_guia"] = int(len(tabela) - len(out))
    out = _limpar_linhas(out)
    out["_tabela"] = nome_tabela
    info["linhas"] = len(out)
    return out, mapa, info


def _com_cabecalho(df):
    if isinstance(df.columns, pd.MultiIndex):
        cab = [" ".join(str(x) for x in t if "Unnamed" not in str(x)) for t in df.columns]
    else:
        cab = ["" if str(c).startswith("Unnamed") else str(c) for c in df.columns]
    return pd.DataFrame([cab] + df.astype(object).values.tolist())


# --------------------------------------------------------------------------
# Análise: duplicidade, consolidação, valor de referência e glosa
# --------------------------------------------------------------------------
_ESPECIALIDADES_POR_TERMO = [
    (r"PSICOMOTRIC|\bPSICOMO?T?\b|\bPSICOM\b", "PSICOMOTRICIDADE"),
    (r"PSICOPEDAG", "PSICOPEDAGOGIA"),
    (r"TERAPIA OCUPAC|TERAPEUTA OCUPAC|\bC TO\b|\bT O\b", "TERAPIA OCUPACIONAL"),
    (r"FONOAUDIOLOG|\bFONO\b", "FONOAUDIOLOGIA"),
    (r"NUTRICION|NUTRICAO|\bNUTRI", "NUTRIÇÃO"),
    (r"FISIOTERAP|\bFISIO\b|\bRPG\b", "FISIOTERAPIA"),
    (r"PSICOTERAPIA|PSICOLOG|\bPSICO\b", "PSICOLOGIA"),
]


def especialidade_do_procedimento(descricao):
    n = _norm(descricao)
    if not n or re.search(r"NEURO|AVALIACAO|AVAL\b", n):
        return None
    for padrao, esp in _ESPECIALIDADES_POR_TERMO:
        if re.search(padrao, n):
            return esp
    return None


def _plano_casa(plano_tabela, operadora):
    a, b = _norm(plano_tabela), _norm(operadora)
    if not a or not b:
        return False
    pa, pb = a.split()[0], b.split()[0]
    return a == b or a in b or b in a or (len(pa) >= 4 and pa == pb)


def data_iso(v):
    """'dd/mm/aaaa', 'aaaa-mm-dd', date ou datetime -> 'aaaa-mm-dd' (ou '' se não for data)."""
    if v is None:
        return ""
    if hasattr(v, "strftime") and not isinstance(v, str):
        try:
            return "" if pd.isna(v) else v.strftime("%Y-%m-%d")
        except (ValueError, TypeError):
            return ""
    s = _txt(v)
    m = _RE_DATA_ISO.match(s)
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    m = _RE_DATA_BR.match(s)
    if m:
        return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    return ""


def vigente(inicio, fim, data):
    """O valor vale na data do atendimento? Sem data: vale o que está em vigor (sem fim)."""
    ini, fi, d = data_iso(inicio), data_iso(fim), data_iso(data)
    if not d:
        return not fi
    return (not ini or ini <= d) and (not fi or fi >= d)


def _linhas_vigencia(df, col_valor):
    ini = df["vigencia_inicio"] if "vigencia_inicio" in df.columns else pd.Series([None] * len(df), index=df.index)
    fim = df["vigencia_fim"] if "vigencia_fim" in df.columns else pd.Series([None] * len(df), index=df.index)
    return [(a, b, float(v)) for a, b, v in zip(ini, fim, df[col_valor]) if pd.notna(v)]


def valor_vigente(linhas, data):
    """linhas: [(inicio, fim, valor)]. Devolve o valor válido na data, ou None (nenhum ou ambíguo)."""
    vals = {round(v, 2) for ini, fim, v in linhas if vigente(ini, fim, data)}
    return vals.pop() if len(vals) == 1 else None


def _ref_cadastro(tabela, operadora, codigo, procedimento, cache, data=""):
    """Valor que a clínica cadastrou para o procedimento nesse plano, na data do atendimento (ou None).
    Procura primeiro pelo código TUSS; se o código não está cadastrado, pela especialidade."""
    if tabela is None or tabela.empty or not operadora:
        return None
    chave = (operadora, codigo, procedimento)
    if chave not in cache:
        t = tabela[tabela["plano"].map(lambda p: _plano_casa(p, operadora))]
        por_cod, por_esp = [], []
        if not t.empty:
            if "codigo_tuss" in t.columns and codigo:
                pc = t[t["codigo_tuss"].fillna("").astype(str).map(_id_txt).str.replace(r"\D", "", regex=True)
                       == re.sub(r"\D", "", str(codigo))]
                por_cod = _linhas_vigencia(pc, "valor_plano")
            esp = especialidade_do_procedimento(procedimento)
            if esp:
                pe = t[t["especialidade"].map(_norm) == _norm(esp)]
                if "codigo_tuss" in pe.columns:  # linhas com código de outro procedimento não valem
                    sem_cod = pe["codigo_tuss"].fillna("").astype(str).str.strip().eq("")
                    pe = pe[sem_cod] if sem_cod.any() else pe
                por_esp = _linhas_vigencia(pe, "valor_plano")
        cache[chave] = (por_cod, por_esp)
    por_cod, por_esp = cache[chave]
    if por_cod:  # código cadastrado: só vale o valor do código (mesmo que não haja vigência na data)
        return valor_vigente(por_cod, data)
    return valor_vigente(por_esp, data) if por_esp else None


def _moda(serie, minimo=3):
    s = serie.round(2)
    s = s[s > 0]
    if len(s) < minimo:
        return None
    contagem = s.value_counts()
    if contagem.iloc[0] < max(minimo, 0.3 * len(s)):
        return None
    return float(contagem.index[0])


def aprender_referencias(df, operadora):
    """Valor que o plano paga normalmente por unidade de cada código: o mais frequente entre os
    itens pagos por inteiro (sem justificativa e com toda a quantidade liberada).
    Retorna {(operadora_norm, codigo, 'AAAA-MM' ou ''): valor}."""
    if df is None or df.empty or not operadora:
        return {}
    d = df[df["Código"].fillna("").astype(str).str.strip().ne("") & df["Código"].ne("-")].copy()
    if d.empty:
        return {}
    qs = pd.to_numeric(d["Quantidade"], errors="coerce")
    ql = pd.to_numeric(d["Quantidade liberada"], errors="coerce")
    unid = qs.where(qs > 0, 1.0).fillna(1.0)
    cheio = ((ql.isna()) | (qs.isna()) | (ql >= qs)) & d["Motivo da glosa"].fillna("").eq("") \
        & d["Código da glosa"].fillna("").eq("")
    d = d[cheio & (pd.to_numeric(d["Valor pago"], errors="coerce") > 0)]
    if d.empty:
        return {}
    d["_unit"] = pd.to_numeric(d["Valor pago"], errors="coerce") / unid[d.index]
    d["_mes"] = pd.to_datetime(d["Data de pagamento"], format="%d/%m/%Y", errors="coerce").dt.strftime("%Y-%m").fillna("")
    op = _norm(operadora)
    refs = {}
    for cod, g in d.groupby("Código"):
        geral = _moda(g["_unit"])
        if geral:
            refs[(op, str(cod), "")] = geral
        for mes, gm in g.groupby("_mes"):
            v = _moda(gm["_unit"])
            if mes and v:
                refs[(op, str(cod), mes)] = v
    return refs


def chave_motivo(motivo):
    """Agrupa variações do mesmo motivo: '1705 VALOR...' e '1705-VALOR...' viram '1705'; textos sem
    código são comparados sem acento/pontuação. Usa só o primeiro motivo quando há vários."""
    primeiro = re.split(r"\s*;\s*", str(motivo or "").strip())[0]
    r = _RE_COD_GLOSA.match(primeiro)
    return r.group(1) if r else _norm(primeiro)


def agrupar_motivos(serie):
    """Para cada motivo, um rótulo único (a forma mais frequente do motivo principal)."""
    s = serie.fillna("").astype(str)
    principal = s.map(lambda m: re.split(r"\s*;\s*", m.strip())[0])
    chave = s.map(chave_motivo)
    rotulo = principal.groupby(chave).agg(lambda x: x.value_counts().index[0])
    return chave.map(rotulo)


def aprender_codigos_glosa(df):
    """Significado dos códigos de glosa vistos nas justificativas ('1705 VALOR APRESENTADO A MAIOR')."""
    res = {}
    if df is None or df.empty:
        return res
    for m in df["Motivo da glosa"].fillna("").astype(str):
        for parte in re.split(r"\s*;\s*", m):
            r = _RE_COD_GLOSA.match(parte)
            if r and len(r.group(2).strip()) >= 6:
                res.setdefault(r.group(1), r.group(2).strip().capitalize())
    return res


def _remover_copias(d):
    """O mesmo pagamento (mesmo atendimento, mesmo crédito, mesmo valor) aparecendo em abas/tabelas
    diferentes é cópia: conta uma vez só. Repetições dentro da MESMA aba são mantidas (podem ser
    várias sessões na mesma guia)."""
    chave_ok = d["Doc Original"].fillna("-").ne("-") & d["Código"].fillna("").ne("") & d["Código"].ne("-")
    if not chave_ok.any() or d["_tabela"].nunique() < 2:
        return d, 0
    assin = ["_k_doc", "_k_cod", "_k_data", "_k_pac", "Data de pagamento", "_k_pago", "_k_qs", "_k_ql"]
    x = d.copy()
    x["_k_doc"], x["_k_cod"], x["_k_data"] = x["Doc Original"], x["Código"], x["Data do evento"]
    x["_k_pac"] = x["Nome do paciente"].map(_norm)
    x["_k_pago"] = pd.to_numeric(x["Valor pago"], errors="coerce").round(2).fillna(-1)
    x["_k_qs"] = pd.to_numeric(x["Quantidade"], errors="coerce").fillna(-1)
    x["_k_ql"] = pd.to_numeric(x["Quantidade liberada"], errors="coerce").fillna(-1)
    x["_ocorr"] = x.groupby(assin + ["_tabela"]).cumcount()
    # para cada assinatura, quantas vezes ela aparece na aba onde aparece mais
    por_aba = x[chave_ok].groupby(assin + ["_tabela"]).size().rename("_n").reset_index()
    maximo = por_aba.groupby(assin)["_n"].max().rename("_max").reset_index()
    melhor_aba = por_aba.merge(maximo, on=assin)
    melhor_aba = melhor_aba[melhor_aba["_n"] == melhor_aba["_max"]].drop_duplicates(assin)[assin + ["_tabela"]]
    melhor_aba = melhor_aba.rename(columns={"_tabela": "_aba_ref"})
    x = x.merge(melhor_aba, on=assin, how="left")
    manter = (~chave_ok.values) | (x["_tabela"] == x["_aba_ref"]) | x["_aba_ref"].isna()
    removidas = int((~manter).sum())
    x = x[manter.values].drop(columns=[c for c in x.columns if c.startswith("_k_")] + ["_ocorr", "_aba_ref"])
    return x.reset_index(drop=True), removidas


def _juntar_texto(serie):
    vs = []
    for v in serie:
        v = str(v).strip() if v is not None and not (isinstance(v, float) and pd.isna(v)) else ""
        if v and v not in ("-", "nan", "None") and v not in vs:
            vs.append(v)
    return "; ".join(vs)


def _ultima_data(serie):
    """Data do último crédito (o pagamento que fecha o atendimento)."""
    dts = pd.to_datetime(pd.Series([x for x in serie if x not in ("", "-")], dtype=object),
                         format="%d/%m/%Y", errors="coerce").dropna()
    return dts.max().strftime("%d/%m/%Y") if len(dts) else "-"


def _consolidar(d):
    """Junta as linhas do mesmo atendimento (mesma guia + código + data + paciente): várias sessões na
    mesma guia, pagamento em mais de um crédito (complemento/recurso)."""
    chave_ok = d["Doc Original"].fillna("-").ne("-") & d["Código"].fillna("").ne("") & d["Código"].ne("-")
    d = d.copy()
    d["_k_pac"] = d["Nome do paciente"].map(_norm)
    d["_grupo"] = range(len(d))
    if chave_ok.any():
        chaves = d.loc[chave_ok, ["Doc Original", "Código", "Data do evento", "_k_pac"]].astype(str).agg("|".join, axis=1)
        d.loc[chave_ok, "_grupo"] = chaves.map({k: i + len(d) for i, k in enumerate(pd.unique(chaves))})
    tamanho = d.groupby("_grupo")["_grupo"].transform("size")
    unicos = d[tamanho == 1].copy()
    unicos["_linhas"] = 1
    unicos["_creditos"] = unicos["Data de pagamento"].map(lambda x: 0 if x in ("", "-") else 1)
    d = d[tamanho > 1]
    if d.empty:
        return unicos.drop(columns=["_grupo", "_k_pac"]).reset_index(drop=True)
    soma = lambda s: float(pd.to_numeric(s, errors="coerce").sum(min_count=1)) if pd.to_numeric(s, errors="coerce").notna().any() else float("nan")
    agg = {
        "Data do evento": "first", "Protocolo": _juntar_texto, "Doc Original": "first",
        "Nome do paciente": "first", "Código": "first", "Procedimento": "first",
        "Valor informado": soma, "Valor pago": soma, "Valor da glosa": soma,
        "Motivo da glosa": _juntar_texto, "Código da glosa": _juntar_texto,
        "Quantidade": soma, "Quantidade liberada": soma, "Valor unitário": "first", "Lote": _juntar_texto,
        "Data de pagamento": _ultima_data,
        "_status_glosa": lambda s: True if any(v is True for v in s) else (False if any(v is False for v in s) else None),
        "_informado_no_arquivo": "max", "_glosa_no_arquivo": "max", "_tabela": "first",
    }
    agg = {k: v for k, v in agg.items() if k in d.columns}
    g = d.groupby("_grupo", sort=False).agg(**{k: (k, v) for k, v in agg.items()})
    g["_linhas"] = d.groupby("_grupo", sort=False).size()
    g["_creditos"] = d.groupby("_grupo", sort=False)["Data de pagamento"].agg(
        lambda s: len({x for x in s if x not in ("", "-")}))
    g["_ordem"] = d.groupby("_grupo", sort=False).apply(lambda x: x.index.min(), include_groups=False)
    unicos["_ordem"] = unicos.index
    for c in ("Protocolo", "Lote", "Doc Original"):
        if c in g.columns:
            g[c] = g[c].replace("", "-")
    junto = pd.concat([unicos.drop(columns=["_grupo", "_k_pac"]), g.reset_index(drop=True)], ignore_index=True)
    return junto.sort_values("_ordem", kind="stable").drop(columns=["_ordem"]).reset_index(drop=True)


def _analisar(d, operadora, conhecimento):
    """Preenche valor informado (quando o arquivo não traz), valor e situação da glosa, motivo e alertas."""
    conhecimento = conhecimento or {}
    tabela = conhecimento.get("tabela_valores")
    refs_hist = conhecimento.get("referencias") or {}
    cod_glosa = dict(conhecimento.get("codigos_glosa") or {})
    refs_arquivo = aprender_referencias(d, operadora)
    cod_glosa.update({k: v for k, v in aprender_codigos_glosa(d).items() if k not in cod_glosa})
    op = _norm(operadora or "")
    cache_cad = {}

    qs = pd.to_numeric(d["Quantidade"], errors="coerce")
    ql = pd.to_numeric(d["Quantidade liberada"], errors="coerce")
    unid = qs.where(qs > 0, 1.0).fillna(1.0)
    pago = pd.to_numeric(d["Valor pago"], errors="coerce").fillna(0.0)
    informado = pd.to_numeric(d["Valor informado"], errors="coerce")
    glosa_arq = pd.to_numeric(d["Valor da glosa"], errors="coerce")
    mes = pd.to_datetime(d["Data de pagamento"].astype(str).str[:10], format="%d/%m/%Y", errors="coerce").dt.strftime("%Y-%m").fillna("")

    ref, fonte = [], []
    for i in range(len(d)):
        cod = _txt(d["Código"].iat[i])
        if bool(d["_informado_no_arquivo"].iat[i]) and pd.notna(informado.iat[i]):
            r = float(informado.iat[i]) / unid.iat[i] if unid.iat[i] else float(informado.iat[i])
            ref.append(round(r, 2)); fonte.append(FONTE_ARQUIVO)
            continue
        r = _ref_cadastro(tabela, operadora, cod, d["Procedimento"].iat[i], cache_cad, d["Data do evento"].iat[i])
        if r is not None:
            ref.append(r); fonte.append(FONTE_CADASTRO)
            continue
        r = (refs_arquivo.get((op, cod, mes.iat[i])) or refs_arquivo.get((op, cod, ""))
             or refs_hist.get((op, cod, mes.iat[i])) or refs_hist.get((op, cod, "")))
        if r:
            ref.append(r); fonte.append(FONTE_APRENDIDO)
        else:
            ref.append(float("nan")); fonte.append(FONTE_NENHUMA)
    ref = pd.Series(ref, index=d.index, dtype="float64")
    sem_inf_arquivo = ~(d["_informado_no_arquivo"].astype(bool) & informado.notna())
    informado = informado.where(~sem_inf_arquivo, (ref * unid).round(2))

    motivo = d["Motivo da glosa"].fillna("").astype(str)
    codg = d["Código da glosa"].fillna("").astype(str)
    # código de glosa sem texto -> texto aprendido / tabela TUSS
    def _texto_motivo(m, c):
        if m.strip():
            return m
        textos = [f"{x} {cod_glosa[x]}" if x in cod_glosa else f"Código de glosa {x}"
                  for x in re.findall(r"\d{4}", c)]
        return "; ".join(textos)
    motivo = pd.Series([_texto_motivo(m, c) for m, c in zip(motivo, codg)], index=d.index)
    tem_just = motivo.str.strip().ne("") | codg.str.strip().ne("") | d["_status_glosa"].map(lambda v: v is True)
    qtd_cortada = ql.notna() & qs.notna() & (ql < qs)

    falta = (informado - pago).round(2)
    glosa = pd.Series(0.0, index=d.index)
    dif = pd.Series(0.0, index=d.index)
    usa_glosa_arquivo = d["_glosa_no_arquivo"].astype(bool) & glosa_arq.notna()
    glosa[usa_glosa_arquivo] = glosa_arq[usa_glosa_arquivo].abs().round(2)
    # arquivo com coluna de glosa vazia/zerada, mas com informado e pago: glosa = informado - pago
    if (d["_glosa_no_arquivo"].astype(bool)).any() and not (glosa_arq.fillna(0) > 0).any():
        calc = d["_informado_no_arquivo"].astype(bool) & informado.notna()
        glosa[calc] = falta[calc].clip(lower=0)
    sem_col = ~d["_glosa_no_arquivo"].astype(bool)
    com_evid = sem_col & (tem_just | qtd_cortada) & informado.notna() & (falta > TOL)
    glosa[com_evid] = falta[com_evid]
    # arquivo traz informado e pago (sem coluna de glosa): a diferença é glosa
    inf_pago_arquivo = sem_col & d["_informado_no_arquivo"].astype(bool) & ~com_evid & (falta > TOL)
    glosa[inf_pago_arquivo] = falta[inf_pago_arquivo]
    # status do arquivo confirma glosa que veio sem valor na linha
    conf_status = d["_status_glosa"].map(lambda v: v is True) & (glosa <= TOL) & (falta > TOL)
    glosa[conf_status] = falta[conf_status]
    sem_evid = sem_col & ~com_evid & ~inf_pago_arquivo & informado.notna() & (falta > TOL)
    dif[sem_evid] = falta[sem_evid]

    sit, alertas = [], []
    for i in range(len(d)):
        g, p, a = glosa.iat[i], pago.iat[i], []
        if g > TOL:
            s = "Glosa total" if p <= TOL else "Glosa parcial"
        elif (tem_just.iat[i] or qtd_cortada.iat[i]) and pd.isna(informado.iat[i]):
            s = "Glosa (valor não identificado)"
            a.append("Sem valor de referência para calcular a glosa: cadastre o valor deste procedimento")
        elif dif.iat[i] > TOL:
            s = "Pago a menor sem justificativa"
            if informado.iat[i] and informado.iat[i] > 0:
                a.append(f"Pago {p / informado.iat[i] * 100:.0f}% do valor de referência".replace(".", ","))
        else:
            s = "Sem glosa"
        if d["_creditos"].iat[i] > 1:
            if tem_just.iat[i] and g <= TOL:
                a.append(f"Glosa revertida: pago em {d['_creditos'].iat[i]} créditos")
            else:
                a.append(f"Pago em {d['_creditos'].iat[i]} créditos")
        if pd.notna(informado.iat[i]) and p > informado.iat[i] * 1.05 + TOL and fonte[i] != FONTE_ARQUIVO:
            a.append("Possível pagamento em duplicidade pelo plano" if d["_creditos"].iat[i] > 1
                     else "Pago acima do valor de referência")
        if tem_just.iat[i] and g <= TOL and dif.iat[i] <= TOL and pd.notna(informado.iat[i]) and fonte[i] == FONTE_APRENDIDO:
            a.append("Plano informou motivo de glosa, mas pagou o valor de referência: cadastre o valor do "
                     "procedimento para calcular a diferença")
        if fonte[i] == FONTE_NENHUMA and not a:
            a.append("Sem valor de referência")
        sit.append(s)
        alertas.append("; ".join(a))

    # motivo para glosa sem justificativa no arquivo
    motivo_final = []
    for i in range(len(d)):
        m = motivo.iat[i]
        if not m.strip() and glosa.iat[i] > TOL:
            m = ("Quantidade liberada menor que a solicitada (sem justificativa no demonstrativo)"
                 if qtd_cortada.iat[i] else
                 "Glosa Silenciosa: Valor zerado ou cortado sem justificativa informada no demonstrativo")
        motivo_final.append(m)

    out = d.copy()
    out["Valor informado"] = informado.round(2)
    out["Valor pago"] = pago.round(2)
    out["Valor da glosa"] = glosa.round(2)
    out["Diferença sem justificativa"] = dif.round(2)
    out["Valor de referência"] = ref.round(2)
    out["Fonte do valor informado"] = fonte
    out["Status da glosa"] = sit
    out["Motivo da glosa"] = motivo_final
    out["Alertas"] = alertas
    out["Quantidade"] = qs
    out["Quantidade liberada"] = ql
    vu = pd.to_numeric(out["Valor unitário"], errors="coerce")
    out["Valor unitário"] = vu.where(vu.notna() & (vu > 0), ref).round(2)
    return out, refs_arquivo, cod_glosa


def finalizar(linhas, operadora, conhecimento=None):
    """Junta as tabelas lidas e faz a análise completa. Retorna (df, resumo)."""
    resumo = dict(linhas_lidas=0, copias_removidas=0, atendimentos=0, linhas_agrupadas=0,
                  referencias={}, codigos_glosa={})
    if not linhas:
        return pd.DataFrame(columns=COLUNAS_COMPLETAS), resumo
    d = pd.concat(linhas, ignore_index=True)
    for c in _MARCADORES:
        if c not in d.columns:
            d[c] = None
    for c in ("Código", "Procedimento", "Motivo da glosa", "Código da glosa"):
        if c in d.columns:
            d[c] = d[c].astype(object).where(d[c].notna(), "").astype(str)
    resumo["linhas_lidas"] = len(d)
    d, resumo["copias_removidas"] = _remover_copias(d)
    g = _consolidar(d)
    resumo["linhas_agrupadas"] = int(len(d) - len(g))
    out, refs, codigos = _analisar(g, operadora, conhecimento)
    resumo.update(atendimentos=len(out), referencias=refs, codigos_glosa=codigos)
    for c in COLUNAS_COMPLETAS:
        if c not in out.columns:
            out[c] = ""
    return out[COLUNAS_COMPLETAS].reset_index(drop=True), resumo


def mapear_dataframe_para_glosas(df_bruto, somente_glosas=True):
    if df_bruto is None or df_bruto.empty:
        return pd.DataFrame(columns=COLUNAS_COMPLETAS)
    linhas, mapa, _ = _processar_bruto(_com_cabecalho(df_bruto))
    df, _ = finalizar([linhas] if not linhas.empty else [], "")
    if somente_glosas:
        df = df[df["Valor da glosa"] > 0.005].reset_index(drop=True)
    df.attrs["mapeamento"] = mapa
    return df


# --------------------------------------------------------------------------
# Leitura de texto corrido (PDF sem tabelas) COM MEMÓRIA DE ESTADO
# --------------------------------------------------------------------------
def _linhas_texto_para_df(linhas):
    regs = []
    contexto = {"data": "-", "nome": "", "protocolo": "-", "guia": "-", "lote": "-", "pagto": "-"}

    for ln in linhas:
        m_prot = _RE_PROTOCOLO.search(ln)
        if m_prot:
            contexto["protocolo"] = m_prot.group(1)

        m_guia = _RE_GUIA.search(ln)
        if m_guia:
            contexto["guia"] = m_guia.group(1)

        m_lote = _RE_LOTE.search(ln)
        if m_lote:
            contexto["lote"] = m_lote.group(1)

        m_data = _RE_DATA.search(ln)
        if m_data:
            if re.search(r"pagamento|pgto|cr[eé]dito", ln, re.I):
                contexto["pagto"] = m_data.group(1)
            else:
                contexto["data"] = m_data.group(1)

        valores = _RE_DIN_TEXTO.findall(ln)
        cod = re.search(r"\b(\d{8}|\d{1,2}\.\d{2}\.\d{2}\.\d{2}-?\d?)\b", ln)

        if not valores and not cod and len(ln) > 6 and _RE_NOME.match(ln):
            palavras = _norm(ln).split()
            if 2 <= len(palavras) <= 6 and not (set(palavras) & _PALAVRAS_PROC):
                contexto["nome"] = ln.strip()

        if len(valores) < 2:
            continue

        resto = _RE_DATA.sub(" ", ln)
        resto = _RE_DIN_TEXTO.sub(" ", resto)
        if cod:
            resto = resto.replace(cod.group(1), " ")
        resto = re.sub(r"\s+", " ", resto).strip(" -|;")

        v = [limpar_valor_monetario(x) for x in valores]
        inf, pago = v[0], v[1]
        glosa = v[2] if len(v) >= 3 else max(inf - pago, 0.0)

        regs.append({
            "Data do evento": padronizar_data(contexto["data"]),
            "Protocolo": contexto["protocolo"] or "-",
            "Doc Original": contexto["guia"] or "-",
            "Nome do paciente": contexto["nome"] or "Não identificado",
            "Código": cod.group(1) if cod else "-",
            "Procedimento": resto,
            "Valor informado": inf, "Valor pago": pago, "Valor da glosa": glosa,
            "Motivo da glosa": "", "Código da glosa": "",
            "Quantidade": float("nan"), "Quantidade liberada": float("nan"),
            "Valor unitário": float("nan"),
            "Lote": contexto["lote"] or "-",
            "Data de pagamento": padronizar_data(contexto["pagto"]),
            "_status_glosa": None, "_informado_no_arquivo": True, "_glosa_no_arquivo": True,
            "_tabela": "texto corrido",
        })
    return pd.DataFrame(regs)


def _linhas_pdf(doc):
    linhas = []
    for pagina in doc:
        palavras = pagina.get_text("words")
        grupos = {}
        for x0, y0, x1, y1, w, *_ in palavras:
            grupos.setdefault(round(y0 / 3), []).append((x0, w))
        for y in sorted(grupos):
            linhas.append(" ".join(w for _, w in sorted(grupos[y])))
    return linhas


# --------------------------------------------------------------------------
# XML TISS (e XMLs parecidos)
# --------------------------------------------------------------------------
def _tag(t):
    return t.split("}")[-1]


def _eh_numero(s):
    return bool(re.fullmatch(r"-?\d+([.,]\d+)?", (s or "").strip()))


def extrair_glosas_xml(arquivo_xml):
    """Lê XML (TISS ou parecido). Cada registro = elemento mais interno que tem dados próprios.
    Subgrupos só de campos simples (ex.: <procedimento>, <relacaoGlosa>) são incorporados ao item:
    valores repetidos são somados (várias glosas no mesmo item) e textos/códigos são juntados."""
    arquivo_xml.seek(0)
    root = ET.parse(arquivo_xml).getroot()
    registros = []

    def folhas(el):
        return {_tag(c.tag): (c.text or "").strip() for c in el if len(c) == 0 and (c.text or "").strip()}

    def eh_grupo_simples(el):
        return len(el) > 0 and all(len(c) == 0 for c in el)

    def percorrer(el, contexto):
        proprias = folhas(el)
        grupos = [c for c in el if eh_grupo_simples(c)]
        complexos = [c for c in el if len(c) > 0 and not eh_grupo_simples(c)]
        # subgrupos simples viram campos do próprio elemento
        if proprias and grupos:
            acumulado = {}
            for gr in grupos:
                for k, v in folhas(gr).items():
                    acumulado.setdefault(k, []).append(v)
            for k, vs in acumulado.items():
                if k in proprias:
                    continue
                if len(vs) > 1 and all(_eh_numero(v) for v in vs) and re.search(r"valor", k, re.I):
                    proprias[k] = f"{sum(limpar_valor_monetario(v) for v in vs):.2f}"
                else:
                    proprias[k] = "; ".join(dict.fromkeys(vs))
            filhos = complexos
        else:
            filhos = [c for c in el if len(c) > 0]
        ctx = {**contexto, **proprias}
        emitiu = False
        for filho in filhos:
            emitiu = percorrer(filho, ctx) or emitiu
        tem_valor = any(re.search(r"valor", k, re.I) for k in proprias)
        if not emitiu and len(proprias) >= 2 and tem_valor:
            registros.append(ctx)
            return True
        return emitiu

    percorrer(root, {})
    if not registros:
        return pd.DataFrame(), {}, dict(tabela="xml", colunas=[], assinatura="", mapa={},
                                        origem_mapa="automático", tipo_documento="", linhas=0,
                                        linhas_guia=0, observacoes=[])
    df = pd.DataFrame(registros)
    return _processar_bruto(_com_cabecalho(df), nome_tabela="xml")


def _detectar_separador(texto):
    linhas = [l for l in texto.splitlines() if l.strip()][:40]
    melhor, melhor_score = ";", -1.0
    for sep in (";", "\t", "|", ","):
        contagens = [len(l) for l in csv.reader(linhas, delimiter=sep)]
        if not contagens:
            continue
        moda = max(set(contagens), key=contagens.count)
        if moda < 2:
            continue
        score = (contagens.count(moda) / len(contagens)) * min(moda, 10)
        if score > melhor_score + 1e-9:
            melhor, melhor_score = sep, score
    return melhor


# Palavras do nome do arquivo -> operadora (mesmos nomes da lista PLANOS_SAUDE do app)
_OPERADORAS_POR_NOME = [
    (("UNIMED",), "Unimed Recife"),
    (("BRADESCO", "BRAD"), "Bradesco Saúde"),
    (("SULAMERICA", "SUL AMERICA", "SULAM"), "SulAmérica"),
    (("AMIL",), "Amil"),
    (("CASSI",), "Cassi"),
    (("GEAP",), "GEAP"),
    (("POSTAL", "CORREIOS"), "Postal Saúde"),
    (("ASSEFAZ",), "Assefaz"),
]


def extrair_nome_operadora(nome_arquivo):
    """Tenta descobrir a operadora pelo nome do arquivo (ex.: 'demonstrativo_unimed_09-2026.xlsx')."""
    n = " " + _norm(nome_arquivo) + " "
    compacto = n.replace(" ", "")
    for chaves, operadora in _OPERADORAS_POR_NOME:
        for c in chaves:
            if (" " + c + " ") in n or (len(c) >= 5 and c.replace(" ", "") in compacto):
                return operadora
    return "INDEFINIDO"


def _detectar_operadora(plano, nome_arquivo):
    if plano and plano != PLANO_AUTOMATICO:
        return plano
    return extrair_nome_operadora(nome_arquivo)


# --------------------------------------------------------------------------
# Função principal chamada pelo app
# --------------------------------------------------------------------------
def extrair_glosas_de_arquivo(arquivo_upload, plano=PLANO_AUTOMATICO, somente_glosas=True,
                              data_pagamento_padrao=None, conhecimento=None):
    """Retorna (DataFrame, mensagem_de_erro_ou_None). Uma linha por atendimento/procedimento.

    data_pagamento_padrao: data usada nas linhas em que o demonstrativo não informa o pagamento.
    conhecimento: dict opcional com perfis, tabela_valores, referencias e codigos_glosa
    (ver motor_conhecimento.carregar)."""
    vazio = pd.DataFrame(columns=COLUNAS_COMPLETAS)
    conhecimento = conhecimento or {}
    perfis = conhecimento.get("perfis") or {}
    try:
        nome_original = getattr(arquivo_upload, "name", "") or ""
        nome = nome_original.lower()
        arquivo_upload.seek(0)
        dados = arquivo_upload.read()
        if not dados:
            return vazio, "O arquivo está vazio."
        resultados = []  # (linhas, mapa, info)

        if nome.endswith((".xlsx", ".xlsm", ".xls")):
            try:
                abas = pd.read_excel(io.BytesIO(dados), sheet_name=None, header=None, dtype=object)
            except ImportError as e:
                return vazio, f"Biblioteca ausente para ler Excel ('openpyxl' para .xlsx, 'xlrd' para .xls): {e}"
            for nome_aba, aba in abas.items():  # TODAS as abas, sem limite
                resultados.append(_processar_bruto(aba, perfis, str(nome_aba)))

        elif nome.endswith((".csv", ".txt", ".tsv")):
            texto = None
            for enc in ("utf-8-sig", "latin-1"):
                try:
                    texto = dados.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            if texto is None:
                return vazio, "Erro de encoding no CSV."
            sep = _detectar_separador(texto)
            linhas = [l for l in csv.reader(io.StringIO(texto), delimiter=sep) if any(c.strip() for c in l)]
            if not linhas:
                return vazio, "O arquivo não tem linhas com conteúdo."
            largura = max(len(l) for l in linhas)
            bruto = pd.DataFrame([l + [None] * (largura - len(l)) for l in linhas], dtype=object)
            resultados.append(_processar_bruto(bruto, perfis, "arquivo"))

        elif nome.endswith((".html", ".htm")):
            for i, t in enumerate(pd.read_html(io.BytesIO(dados))):
                resultados.append(_processar_bruto(_com_cabecalho(t), perfis, f"tabela {i + 1}"))

        elif nome.endswith(".xml"):
            resultados.append(extrair_glosas_xml(io.BytesIO(dados)))

        elif nome.endswith(".pdf"):
            if fitz is None:
                return vazio, "PyMuPDF não instalado. Rode: pip install pymupdf"
            doc = fitz.open(stream=dados, filetype="pdf")
            try:
                for n_pag, pagina in enumerate(doc, start=1):
                    if hasattr(pagina, "find_tables"):
                        for t in pagina.find_tables().tables:
                            resultados.append(_processar_bruto(pd.DataFrame(t.extract()), perfis, f"página {n_pag}"))
                if not any(not d.empty for d, _, _ in resultados):
                    texto_df = _limpar_linhas(_linhas_texto_para_df(_linhas_pdf(doc)))
                    resultados.append((texto_df, {"origem": "texto corrido"},
                                       dict(tabela="texto corrido", colunas=[], assinatura="", mapa={},
                                            origem_mapa="texto corrido", tipo_documento="", linhas=len(texto_df),
                                            linhas_guia=0, observacoes=["PDF sem tabela: lido como texto corrido."])))
            finally:
                doc.close()

        elif nome.endswith(".docx"):
            if docx is None:
                return vazio, "python-docx não instalado. Rode: pip install python-docx"
            d = docx.Document(io.BytesIO(dados))
            for i, tab in enumerate(d.tables):
                linhas = [[c.text.strip() for c in lin.cells] for lin in tab.rows]
                resultados.append(_processar_bruto(pd.DataFrame(linhas), perfis, f"tabela {i + 1}"))
            if not any(not x.empty for x, _, _ in resultados):
                linhas = [p.text for p in d.paragraphs if p.text.strip()]
                texto_df = _limpar_linhas(_linhas_texto_para_df(linhas))
                resultados.append((texto_df, {"origem": "texto corrido"},
                                   dict(tabela="texto corrido", colunas=[], assinatura="", mapa={},
                                        origem_mapa="texto corrido", tipo_documento="", linhas=len(texto_df),
                                        linhas_guia=0, observacoes=[])))
        else:
            return vazio, f"Formato não suportado: {nome_original}"

        partes = [d for d, _, _ in resultados if not d.empty]
        mapa = next((m for d, m, _ in resultados if not d.empty), {})
        leituras = [i for d, _, i in resultados if i.get("colunas") or not d.empty]
        if not partes:
            return vazio, "Nenhuma glosa ou estrutura de valores encontrada."

        operadora = _detectar_operadora(plano, nome_original)
        if data_pagamento_padrao is not None and str(data_pagamento_padrao).strip():
            padrao = padronizar_data(data_pagamento_padrao)
            if padrao != "-":
                for p in partes:
                    p.loc[p["Data de pagamento"].isin(["", "-"]), "Data de pagamento"] = padrao

        df, resumo = finalizar(partes, operadora, conhecimento)
        total_atend = len(df)
        if somente_glosas:
            df = df[df["Valor da glosa"] > 0.005].reset_index(drop=True)
        df.attrs["mapeamento"] = mapa
        df.attrs["operadora"] = operadora
        df.attrs["abas_total"] = len(resultados)
        df.attrs["abas_com_dados"] = len(partes)
        df.attrs["linhas_lidas"] = resumo["linhas_lidas"]
        df.attrs["leituras"] = leituras
        df.attrs["resumo"] = {k: v for k, v in resumo.items() if k not in ("referencias", "codigos_glosa")}
        df.attrs["referencias_aprendidas"] = resumo["referencias"]
        df.attrs["codigos_glosa_aprendidos"] = resumo["codigos_glosa"]

        if df.empty:
            return df, (f"Arquivo lido ({total_atend} atendimentos), mas nenhuma glosa maior que zero foi "
                        f"encontrada. Colunas reconhecidas: {mapa}")
        return df, None

    except Exception as e:
        return vazio, f"Erro ao processar o arquivo: {type(e).__name__}: {e}"
