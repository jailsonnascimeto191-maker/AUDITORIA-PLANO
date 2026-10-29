"""Anomalias 'Para verificar', lista de recursos e relatório Excel com as abas novas."""
import io
import unittest

import pandas as pd
from openpyxl import load_workbook

from base import CAB, analisado, ler, linha, me
import motor_padroes as mp
import motor_relatorios as mr
import motor_verificacao as mv

MOT = "1705 VALOR APRESENTADO A MAIOR"


def _r(**kw):
    base = {"Data do evento": "05/01/2025", "Nome do paciente": "ANA TESTE SILVA", "Código": "50000470",
            "Procedimento": "SESSAO", "Doc Original": "G1", "Protocolo": "1100000001", "Valor informado": 100.0,
            "Valor pago": 100.0, "Data de pagamento": "10/02/2025", "Status da glosa": "Sem glosa"}
    base.update(kw)
    return base


class TestAnomalias(unittest.TestCase):
    def test_vazio(self):
        a = mv.anomalias(pd.DataFrame())
        self.assertTrue(a.empty)
        self.assertEqual(list(a.columns), mv.COLUNAS_ANOMALIA)

    def test_sem_nada_estranho(self):
        self.assertTrue(mv.anomalias(analisado([_r()])).empty)

    def test_pago_repetido(self):
        a = mv.anomalias(analisado([_r(), _r(**{"Status da glosa": me.STATUS_DUPLICIDADE, "Valor informado": 0.0,
                                                "Data de pagamento": "10/04/2025"})]))
        self.assertEqual(list(a["Tipo"]), [mv.TIPO_PAGO_REPETIDO])

    def test_cobranca_em_dobro(self):
        df = analisado([_r(**{"Valor pago": 0.0, "Situação do recurso": me.RECURSO_RECEBIDO, "Glosa original": 100.0,
                              "Recebido no recurso": 100.0, "Motivo da glosa": "1806 DUPLICIDADE DE COBRANCA",
                              "Status da glosa": me.STATUS_RECURSO_PAGO, "Complemento recebido": 100.0}),
                        _r(**{"Doc Original": "G9"})])
        # o atendimento original tem valor pago 0 (o recurso veio em outra linha): ainda assim é caso para verificar
        a = mv.anomalias(df)
        self.assertIn(mv.TIPO_COBRANCA_DOBRO, list(a["Tipo"]))

    def test_cobranca_em_dobro_original_pago(self):
        df = analisado([_r(**{"Valor pago": 100.0, "Situação do recurso": me.RECURSO_RECEBIDO, "Glosa original": 100.0,
                              "Recebido no recurso": 100.0, "Status da glosa": me.STATUS_RECURSO_PAGO}),
                        _r(**{"Doc Original": "G9"})])
        a = mv.anomalias(df)
        self.assertEqual(list(a["Tipo"]), [mv.TIPO_COBRANCA_DOBRO])
        self.assertIn("G9", a["O que verificar"].iloc[0])

    def test_acima_e_varios_originais(self):
        df = analisado([_r(Alertas=mp.ALERTA_ACIMA, **{"Complemento recebido": 60.0}),
                        _r(Alertas=mp.ALERTA_VARIOS_ORIGINAIS, **{"Doc Original": "G2", "Código": "1"})])
        self.assertEqual(sorted(mv.anomalias(df)["Tipo"]), sorted([mv.TIPO_ACIMA, mv.TIPO_SEM_ORIGINAL_DEFINIDO]))

    def test_colunas_minimas_sem_protocolo(self):
        df = analisado([_r(**{"Status da glosa": me.STATUS_DUPLICIDADE})]).drop(columns=["Protocolo"])
        try:
            a = mv.anomalias(df)
        except Exception as e:
            self.fail(f"{type(e).__name__}: {e}")
        self.assertEqual(len(a), 1)

    def test_recebeu_mais_que_a_glosa_no_recurso_vira_caso(self):
        # glosa 100, recurso pagou 100 e o plano pagou outra vez (200 recebidos para 100 informados)
        um, _ = ler({"a": [CAB, linha(pago=0, glosa=100, motivo=MOT, dtpag="10/02/2025"),
                           linha(pago=100, glosa=0, dtpag="10/04/2025", prot="2200000001"),
                           linha(pago=100, glosa=0, dtpag="10/06/2025", prot="2200000003")]})
        self.assertGreater(len(mv.anomalias(um)), 0, um[["Valor pago", "Status da glosa", "Alertas"]].to_string())


class TestRecursos(unittest.TestCase):
    def test_lista_recursos(self):
        df, _ = ler({"a": [CAB, linha(pago=0, glosa=100, motivo=MOT, dtpag="10/02/2025"),
                           linha(pago=40, glosa=60, motivo=MOT, dtpag="10/04/2025", prot="2200000001")]})
        r = mv.recursos(df)
        self.assertEqual(len(r), 1)
        x = r.iloc[0]
        self.assertEqual((x["Valor glosado"], x["Valor recebido"], x["Falta receber"], x["Situação"]),
                         (100.0, 40.0, 60.0, me.RECURSO_PARCIAL))
        self.assertEqual(x["Protocolo da glosa"], "1100000001")
        self.assertEqual(x["Protocolo do recurso"], "2200000001")

    def test_sem_recursos(self):
        self.assertTrue(mv.recursos(analisado([_r()])).empty)


class TestRelatorioExcel(unittest.TestCase):
    def _gerar(self, linhas):
        df, _ = ler({"a": [CAB] + linhas})
        gl = df[df["Valor da glosa"] > 0.005]
        spec = mr.spec_glosas(gl, df, ["demo.xlsx"])
        return spec, load_workbook(io.BytesIO(mr.gerar_excel(spec)))

    def test_abas_novas_com_dados(self):
        spec, wb = self._gerar([linha(pago=0, glosa=100, motivo=MOT, dtpag="10/02/2025"),
                                linha(pago=40, glosa=60, motivo=MOT, dtpag="10/04/2025", prot="2200000001"),
                                linha(guia="900002", pac="BRUNO TESTE LIMA", pago=60, glosa=40, motivo=MOT)])
        self.assertEqual(wb.sheetnames, ["Resumo", "Glosas", "Recursos de glosa", "Para verificar"])
        ws = wb["Recursos de glosa"]
        valores = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
        self.assertIn(me.RECURSO_PARCIAL, valores)
        kpis = dict((k, v) for k, v, _ in spec["kpis"])
        self.assertEqual(kpis["Glosas recursadas"], 1)
        self.assertAlmostEqual(kpis["Ainda a receber dos recursos"], 60.0)
        self.assertAlmostEqual(kpis["Valor glosado"], 100.0)

    def test_abas_novas_vazias(self):
        _, wb = self._gerar([linha(pago=60, glosa=40, motivo=MOT)])
        self.assertIn("Para verificar", wb.sheetnames)
        self.assertEqual(wb["Para verificar"]["A4"].value, "Nenhum registro.")

    def test_pdf_nao_quebra(self):
        spec, _ = self._gerar([linha(pago=0, glosa=100, motivo=MOT, dtpag="10/02/2025"),
                               linha(pago=40, glosa=60, motivo=MOT, dtpag="10/04/2025", prot="2200000001")])
        pdf = mr.gerar_pdf(spec)
        self.assertTrue(pdf.startswith(b"%PDF"))

    def test_relatorio_sem_glosa(self):
        df, _ = ler({"a": [CAB, linha()]})
        spec = mr.spec_glosas(df[df["Valor da glosa"] > 0.005], df, ["demo.xlsx"])
        load_workbook(io.BytesIO(mr.gerar_excel(spec)))
        self.assertTrue(mr.gerar_pdf(spec).startswith(b"%PDF"))

    def test_falha_nas_abas_novas_nao_e_silenciosa(self):
        # se recursos()/anomalias() quebrarem, o relatório sai sem as abas e sem aviso nenhum
        df, _ = ler({"a": [CAB, linha(pago=0, glosa=100, motivo=MOT, dtpag="10/02/2025"),
                           linha(pago=40, glosa=60, motivo=MOT, dtpag="10/04/2025", prot="2200000001")]})
        quebrado = df.drop(columns=["Nome do paciente"])
        spec = mr.spec_glosas(df[df["Valor da glosa"] > 0], quebrado, ["demo.xlsx"])
        textos = " ".join(str(x) for x in spec.get("info", []))
        aba = next(d for d in spec["detalhes"] if d.get("nome") == "Recursos de glosa")
        self.assertTrue(len(aba["df"]) > 0 or "não foi possível" in textos.lower(),
                        "abas de recursos/verificação somem em silêncio quando há erro")


if __name__ == "__main__":
    unittest.main(verbosity=2)
