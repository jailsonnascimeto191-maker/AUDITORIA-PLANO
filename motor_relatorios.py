"""
Motor de relatórios do Espaço Saúde: gera Excel (.xlsx) e PDF a partir de uma "especificação" (spec).

Uso:
    spec = spec_producao(...)  /  spec_glosas(...)  /  spec_faturamento(...)
    xlsx_bytes = gerar_excel(spec)
    pdf_bytes  = gerar_pdf(spec)

A spec é um dicionário simples:
    titulo, subtitulo
    info:      [(rótulo, valor), ...]                      dados de cabeçalho (profissional, mês, ...)
    kpis:      [(rótulo, valor, tipo), ...]                tipo: "moeda" | "int" | "pct" | "texto"
    resumos:   [dict(titulo, df, moeda=[...], inteiros=[...], pct=[...]), ...]
    detalhes:  [dict(nome, titulo, df, moeda=[...], inteiros=[...], datas=[...], total=[...],
                     status="Situação", pdf_colunas=[...], pdf=True), ...]
"""
import io
import re
from datetime import datetime
from xml.sax.saxutils import escape

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

import motor_cruzamento as mc
import motor_extracao as me

# Paleta (mesma do app)
NAVY, TEAL, TEAL_CLARO, CINZA, LINHA = "0F2A43", "0E7C86", "E3F1F2", "F3F5F8", "D5DBE3"
VERDE, AMBAR, VERMELHO = "E6F4EA", "FFF4D6", "FDE7E7"
FONTE = "Arial"
FMT_MOEDA = '"R$" #,##0.00'
MAX_LINHAS_PDF = 4000


# --------------------------------------------------------------------------
# Utilidades
# --------------------------------------------------------------------------
def brl(v):
    try:
        if v is None or pd.isna(v):
            return "-"
        return "R$ " + f"{float(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except (TypeError, ValueError):
        return str(v)


def _agora():
    return datetime.now().strftime("%d/%m/%Y %H:%M")


def _py(v):
    """Converte para tipo Python simples; NaN vira None."""
    try:
        if v is None or (not isinstance(v, str) and pd.isna(v)):
            return None
    except (TypeError, ValueError):
        pass
    if hasattr(v, "item"):
        try:
            return v.item()
        except (ValueError, AttributeError):
            pass
    return v


def _cor_status(txt):
    n = str(txt or "").lower()
    if not n:
        return None
    if "parcial" in n or "com glosa" in n or "diverg" in n or "a menor" in n:
        return AMBAR
    if "glos" in n or "não pago" in n or "nao pago" in n or "não localizado" in n:
        return VERMELHO
    if n.startswith("pago"):
        return VERDE
    return None


def _fmt_pdf(v, tipo):
    v = _py(v)
    if v is None:
        return "-"
    if tipo == "moeda":
        return brl(v)
    if tipo == "int":
        try:
            return f"{float(v):,.0f}".replace(",", ".")
        except (TypeError, ValueError):
            return str(v)
    if tipo == "pct":
        try:
            return f"{float(v) * 100:.1f}%".replace(".", ",")
        except (TypeError, ValueError):
            return str(v)
    return str(v)


def _tipo_coluna(nome, d):
    if nome in d.get("moeda", []):
        return "moeda"
    if nome in d.get("inteiros", []):
        return "int"
    if nome in d.get("pct", []):
        return "pct"
    return "texto"


def _nome_aba(nome, usados):
    base = re.sub(r"[\\/*?:\[\]]", " ", nome).strip()[:31] or "Aba"
    n, i = base, 2
    while n in usados:
        suf = f" ({i})"
        n = base[:31 - len(suf)] + suf
        i += 1
    usados.add(n)
    return n


# --------------------------------------------------------------------------
# EXCEL
# --------------------------------------------------------------------------
def _estilo_cabecalho(cel):
    cel.font = Font(name=FONTE, bold=True, color="FFFFFF", size=10)
    cel.fill = PatternFill("solid", fgColor=NAVY)
    cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


_BORDA = Border(bottom=Side(style="thin", color=LINHA))


def _titulo_aba(ws, titulo, subtitulo, ncols):
    ncols = max(ncols, 4)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=titulo)
    c.font = Font(name=FONTE, bold=True, size=15, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    for col in range(2, ncols + 1):
        ws.cell(row=1, column=col).fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    s = ws.cell(row=2, column=1, value=subtitulo)
    s.font = Font(name=FONTE, size=9, color="637083", italic=True)
    s.alignment = Alignment(indent=1)


def _configurar_impressao(ws, linha_cab=None):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    if linha_cab:
        ws.print_title_rows = f"{linha_cab}:{linha_cab}"
    ws.oddFooter.center.text = "Espaço Saúde  |  Página &P de &N"


def _escrever_tabela(ws, df, linha_ini, d, formatar_status=True, totais=True):
    """Escreve df a partir de linha_ini (cabeçalho). Retorna a última linha usada."""
    cols = list(df.columns)
    datas = set(d.get("datas", []))
    for j, nome in enumerate(cols, start=1):
        _estilo_cabecalho(ws.cell(row=linha_ini, column=j, value=nome))
    ws.row_dimensions[linha_ini].height = 30

    status_col = d.get("status")
    fim = linha_ini
    for i, linha in enumerate(df.itertuples(index=False), start=1):
        r = linha_ini + i
        fim = r
        cor_status = None
        for j, (nome, v) in enumerate(zip(cols, linha), start=1):
            v = _py(v)
            if nome in datas and isinstance(v, str):
                dt = pd.to_datetime(v, format="%d/%m/%Y", errors="coerce")
                v = dt.to_pydatetime() if pd.notna(dt) else (v if v != "-" else None)
            cel = ws.cell(row=r, column=j, value=v)
            cel.font = Font(name=FONTE, size=9)
            cel.border = _BORDA
            tipo = _tipo_coluna(nome, d)
            if tipo == "moeda":
                cel.number_format = FMT_MOEDA
            elif tipo == "int":
                cel.number_format = "#,##0"
            elif tipo == "pct":
                cel.number_format = "0.0%"
            elif nome in datas:
                cel.number_format = "DD/MM/YYYY"
                cel.alignment = Alignment(horizontal="center")
            else:
                cel.alignment = Alignment(vertical="top", wrap_text=(isinstance(v, str) and len(v) > 40))
            if formatar_status and status_col and nome == status_col:
                cor_status = _cor_status(v)
                if cor_status:
                    cel.fill = PatternFill("solid", fgColor=cor_status)
                    cel.font = Font(name=FONTE, size=9, bold=True)
        if i % 2 == 0:
            for j, nome in enumerate(cols, start=1):
                cel = ws.cell(row=r, column=j)
                if cel.fill.fgColor.rgb in (None, "00000000"):
                    cel.fill = PatternFill("solid", fgColor=CINZA)

    # larguras
    for j, nome in enumerate(cols, start=1):
        amostra = [len(str(nome))] + [len(str(_py(x))) for x in df[nome].head(300) if _py(x) is not None]
        ws.column_dimensions[get_column_letter(j)].width = min(max(max(amostra) + 2, 10), 46)

    # totais com SUBTOTAL (respeitam o filtro aplicado pelo usuário)
    if totais and d.get("total") and fim > linha_ini:
        rt = fim + 1
        ws.cell(row=rt, column=1, value="TOTAL").font = Font(name=FONTE, bold=True, size=10, color="FFFFFF")
        for j in range(1, len(cols) + 1):
            ws.cell(row=rt, column=j).fill = PatternFill("solid", fgColor=TEAL)
        for nome in d["total"]:
            if nome in cols:
                j = cols.index(nome) + 1
                L = get_column_letter(j)
                cel = ws.cell(row=rt, column=j, value=f"=SUBTOTAL(109,{L}{linha_ini + 1}:{L}{fim})")
                cel.font = Font(name=FONTE, bold=True, size=10, color="FFFFFF")
                cel.number_format = FMT_MOEDA if _tipo_coluna(nome, d) == "moeda" else "#,##0"
        fim = rt
    return fim


def gerar_excel(spec):
    wb = Workbook()
    usados = set()

    # ---- Aba Resumo ----
    ws = wb.active
    ws.title = _nome_aba("Resumo", usados)
    ws.sheet_view.showGridLines = False
    larg = 6
    _titulo_aba(ws, spec["titulo"], spec.get("subtitulo", "") + f"   |   Emitido em {_agora()}", larg)
    r = 4
    for rot, val in spec.get("info", []):
        a = ws.cell(row=r, column=1, value=rot)
        a.font = Font(name=FONTE, bold=True, size=10, color=NAVY)
        b = ws.cell(row=r, column=2, value=_py(val))
        b.font = Font(name=FONTE, size=10)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=larg)
        r += 1
    r += 1
    if spec.get("kpis"):
        for j, t in enumerate(["Indicador", "Valor"], start=1):
            _estilo_cabecalho(ws.cell(row=r, column=j, value=t))
        r += 1
        for rot, val, tipo in spec["kpis"]:
            a = ws.cell(row=r, column=1, value=rot)
            a.font = Font(name=FONTE, size=10)
            a.fill = PatternFill("solid", fgColor=TEAL_CLARO)
            b = ws.cell(row=r, column=2, value=_py(val))
            b.font = Font(name=FONTE, bold=True, size=11, color=NAVY)
            b.fill = PatternFill("solid", fgColor=TEAL_CLARO)
            b.alignment = Alignment(horizontal="right")
            b.number_format = {"moeda": FMT_MOEDA, "int": "#,##0", "pct": "0.0%"}.get(tipo, "General")
            r += 1
        r += 1
    maxcols = 2
    for res in spec.get("resumos", []):
        df = res["df"]
        if df is None or df.empty:
            continue
        t = ws.cell(row=r, column=1, value=res["titulo"])
        t.font = Font(name=FONTE, bold=True, size=11, color=TEAL)
        r += 1
        fim = _escrever_tabela(ws, df, r, res, formatar_status=False, totais=False)
        maxcols = max(maxcols, df.shape[1])
        r = fim + 2
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 34)
    for j in range(2, maxcols + 1):
        L = get_column_letter(j)
        ws.column_dimensions[L].width = max(ws.column_dimensions[L].width or 0, 18)
    _configurar_impressao(ws)

    # ---- Abas de detalhe ----
    for d in spec.get("detalhes", []):
        df = d["df"]
        ws = wb.create_sheet(_nome_aba(d.get("nome", d["titulo"]), usados))
        ws.sheet_view.showGridLines = False
        n = max(df.shape[1], 4)
        _titulo_aba(ws, f"{spec['titulo']}  -  {d['titulo']}", spec.get("subtitulo", ""), n)
        if df.empty:
            ws.cell(row=4, column=1, value="Nenhum registro.").font = Font(name=FONTE, italic=True, size=10)
            continue
        fim = _escrever_tabela(ws, df, 4, d)
        ult = fim - 1 if d.get("total") else fim
        ws.auto_filter.ref = f"A4:{get_column_letter(df.shape[1])}{ult}"
        ws.freeze_panes = "A5"
        _configurar_impressao(ws, linha_cab=4)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def gerar_pdf(spec):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)

    c_navy, c_teal = colors.HexColor("#" + NAVY), colors.HexColor("#" + TEAL)
    c_cinza, c_linha = colors.HexColor("#" + CINZA), colors.HexColor("#" + LINHA)

    pw, ph = landscape(A4)
    margem = 12 * mm
    util = pw - 2 * margem

    st_celula = ParagraphStyle("c", fontName="Helvetica", fontSize=7, leading=8.5)
    st_cab = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=7.5, leading=9, textColor=colors.white,
                            alignment=1)
    st_sec = ParagraphStyle("s", fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=c_teal,
                            spaceBefore=10, spaceAfter=4)
    st_nota = ParagraphStyle("n", fontName="Helvetica-Oblique", fontSize=7.5, textColor=colors.HexColor("#637083"))
    st_dir = ParagraphStyle("d", parent=st_celula, alignment=TA_RIGHT)

    story = []

    # Faixa de título
    t_titulo = Table(
        [[Paragraph(f"<font color='white' size='16'><b>{escape(spec['titulo'])}</b></font>",
                    ParagraphStyle("t", leading=20))],
         [Paragraph(f"<font color='#C9D6E2' size='8.5'>{escape(spec.get('subtitulo', ''))}</font>",
                    ParagraphStyle("t2", leading=11))]],
        colWidths=[util])
    t_titulo.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), c_navy), ("LEFTPADDING", (0, 0), (-1, -1), 10),
        ("TOPPADDING", (0, 0), (0, 0), 9), ("BOTTOMPADDING", (0, 1), (0, 1), 9),
        ("TOPPADDING", (0, 1), (0, 1), 0), ("BOTTOMPADDING", (0, 0), (0, 0), 2)]))
    story += [t_titulo, Spacer(1, 6)]

    # Dados de cabeçalho (profissional, mês...)
    info = spec.get("info", [])
    if info:
        linhas = []
        for i in range(0, len(info), 2):
            par = info[i:i + 2]
            linha = []
            for rot, val in par:
                linha += [Paragraph(f"<b>{escape(str(rot))}</b>", st_celula),
                          Paragraph(escape(str(_py(val) if _py(val) is not None else '-')), st_celula)]
            while len(linha) < 4:
                linha += ["", ""]
            linhas.append(linha)
        t = Table(linhas, colWidths=[util * 0.14, util * 0.36, util * 0.14, util * 0.36])
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), c_cinza), ("BOX", (0, 0), (-1, -1), 0.5, c_linha),
                               ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
        story += [t, Spacer(1, 8)]

    # Indicadores
    kp = spec.get("kpis", [])
    if kp:
        por_linha = 4
        blocos = []
        for i in range(0, len(kp), por_linha):
            grupo = kp[i:i + por_linha]
            cel = []
            for rot, val, tipo in grupo:
                cel.append([Paragraph(f"<font size='7' color='#637083'>{escape(rot)}</font>", st_celula),
                            Paragraph(f"<font size='12' color='#{NAVY}'><b>{escape(_fmt_pdf(val, tipo))}</b></font>",
                                      ParagraphStyle("k", leading=15))])
            while len(cel) < por_linha:
                cel.append("")
            blocos.append(cel)
        t = Table(blocos, colWidths=[util / por_linha] * por_linha)
        t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.5, c_linha), ("INNERGRID", (0, 0), (-1, -1), 0.5, c_linha),
                               ("LINEABOVE", (0, 0), (-1, 0), 2, c_teal),
                               ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                               ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                               ("LEFTPADDING", (0, 0), (-1, -1), 8), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        story += [t, Spacer(1, 4)]

    def tabela_pdf(titulo, df, d, truncar=True):
        cols = list(d.get("pdf_colunas") or df.columns)
        cols = [c for c in cols if c in df.columns]
        sub = df[cols]
        nota = None
        if truncar and len(sub) > MAX_LINHAS_PDF:
            nota = (f"Exibindo as primeiras {MAX_LINHAS_PDF:,} de {len(sub):,} linhas neste PDF. "
                    "A lista completa está no arquivo Excel.").replace(",", ".")
            sub = sub.head(MAX_LINHAS_PDF)
        tipos = {c: _tipo_coluna(c, d) for c in cols}
        # larguras proporcionais ao conteúdo
        pesos = []
        for c in cols:
            ls = [len(str(c)) * 0.8] + [len(_fmt_pdf(x, tipos[c])) for x in sub[c].head(200)]
            pesos.append(min(max(max(ls), 11), 28 if tipos[c] == "texto" else 16))
        larg = [util * p / sum(pesos) for p in pesos]
        dados = [[Paragraph(escape(str(c)), st_cab) for c in cols]]
        estilos = [("BACKGROUND", (0, 0), (-1, 0), c_navy), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                   ("LINEBELOW", (0, 0), (-1, -1), 0.25, c_linha),
                   ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
                   ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]
        sc = d.get("status")
        for i, linha in enumerate(sub.itertuples(index=False), start=1):
            row = []
            for j, (c, v) in enumerate(zip(cols, linha)):
                txt = _fmt_pdf(v, tipos[c])
                if tipos[c] == "texto" and len(txt) > 14:
                    row.append(Paragraph(escape(txt), st_celula))
                elif tipos[c] in ("moeda", "int", "pct"):
                    row.append(Paragraph(escape(txt), st_dir))
                else:
                    row.append(Paragraph(escape(txt), st_celula))
                if sc and c == sc:
                    cor = _cor_status(_py(v))
                    if cor:
                        estilos.append(("BACKGROUND", (j, i), (j, i), colors.HexColor("#" + cor)))
            if i % 2 == 0:
                estilos.append(("BACKGROUND", (0, i), (-1, i), c_cinza))
            dados.append(row)
        # linha de totais
        if d.get("total"):
            tot = []
            for c in cols:
                if c in d["total"]:
                    tot.append(Paragraph(f"<b>{escape(_fmt_pdf(pd.to_numeric(sub[c], errors='coerce').sum(), tipos[c]))}</b>", st_dir))
                else:
                    tot.append(Paragraph("<b>TOTAL</b>" if c == cols[0] else "", st_celula))
            dados.append(tot)
            estilos.append(("BACKGROUND", (0, len(dados) - 1), (-1, len(dados) - 1), colors.HexColor("#" + TEAL_CLARO)))
        t = Table(dados, colWidths=larg, repeatRows=1)
        t.setStyle(TableStyle(estilos))
        out = [Paragraph(escape(titulo), st_sec), t]
        if nota:
            out += [Spacer(1, 3), Paragraph(nota, st_nota)]
        return out

    for res in spec.get("resumos", []):
        if res["df"] is not None and not res["df"].empty:
            story += tabela_pdf(res["titulo"], res["df"], res, truncar=False)

    for d in spec.get("detalhes", []):
        if not d.get("pdf", True):
            continue
        if d["df"].empty:
            story += [Paragraph(escape(d["titulo"]), st_sec), Paragraph("Nenhum registro.", st_nota)]
        else:
            story += tabela_pdf(d["titulo"], d["df"], d)

    def rodape(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#637083"))
        canvas.drawString(margem, 7 * mm, f"Espaço Saúde  |  {spec['titulo']}  |  Emitido em {_agora()}")
        canvas.drawRightString(pw - margem, 7 * mm, f"Página {doc.page}")
        canvas.setStrokeColor(c_linha)
        canvas.line(margem, 10 * mm, pw - margem, 10 * mm)
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=margem, rightMargin=margem,
                            topMargin=margem, bottomMargin=15 * mm, title=spec["titulo"],
                            author="Espaço Saúde")
    doc.build(story, onFirstPage=rodape, onLaterPages=rodape)
    return buf.getvalue()


# --------------------------------------------------------------------------
# Especificações prontas dos três relatórios
# --------------------------------------------------------------------------
def spec_producao(profissional, registro, especialidade, competencia, espelho, resumo, por_plano,
                  so_demonstrativo=None, descontou_glosa=False):
    """Relatório 1: produção do profissional (espelho de tudo que ele atendeu/recebeu)."""
    moeda = ["Valor unitário (plano)", "Valor total a receber", "Valor pago pelo plano", "Valor da glosa",
             "Repasse unitário", "Repasse ao profissional"]
    nao_loc = espelho[espelho["Situação"].str.startswith("Não localizado")][
        ["Data", "Paciente", "Plano", "Tempo", "Quantidade", "Guia", "Lote"]]
    detalhes = [
        dict(nome="Espelho da produção", titulo="Espelho da produção", df=espelho, moeda=moeda,
             inteiros=["Quantidade"], datas=["Data"], total=["Quantidade", "Valor total a receber",
                                                              "Valor pago pelo plano", "Valor da glosa",
                                                              "Repasse ao profissional"],
             status="Situação",
             pdf_colunas=["Data", "Paciente", "Plano", "Quantidade", "Valor unitário (plano)",
                          "Valor total a receber", "Valor pago pelo plano", "Valor da glosa",
                          "Repasse ao profissional", "Situação", "Motivo da glosa"]),
        dict(nome="Não localizados", titulo="Atendimentos não localizados no demonstrativo", df=nao_loc,
             inteiros=["Quantidade"], datas=["Data"], pdf=True),
    ]
    if so_demonstrativo is not None and not so_demonstrativo.empty:
        detalhes.append(dict(nome="Só no demonstrativo", titulo="Informados pelo plano e ausentes do controle",
                             df=so_demonstrativo, moeda=["Valor informado", "Valor pago", "Valor da glosa"],
                             datas=["Data"], pdf=False))
    info = [("Profissional", profissional), ("Registro", registro or "-"), ("Especialidade", especialidade),
            ("Mês de produção", competencia)]
    kpis = [("Atendimentos", resumo["atendimentos"], "int"), ("Pacientes atendidos", resumo["pacientes"], "int"),
            ("Valor total a receber (tabela)", resumo["receber"], "moeda"),
            ("Valor pago pelo plano", resumo["pago"], "moeda"),
            ("Valor glosado", resumo["glosa"], "moeda"),
            ("Repasse ao profissional" + (" (após glosa)" if descontou_glosa else ""), resumo["repasse"], "moeda"),
            ("Não localizados no demonstrativo", resumo["nao_localizados"], "int")]
    return dict(
        titulo="Relatório de Produção do Profissional",
        subtitulo=f"{profissional}  |  {especialidade}  |  Competência {competencia}",
        info=info, kpis=kpis,
        resumos=[dict(titulo="Resumo por plano de saúde", df=por_plano,
                      moeda=["Valor total a receber", "Valor pago pelo plano", "Valor da glosa",
                             "Repasse ao profissional"], inteiros=["Pacientes", "Atendimentos"])],
        detalhes=detalhes)


def spec_glosas(df_glosas, df_completo=None, arquivos=None):
    """Relatório 2: glosas (data, paciente, guia do prestador, lote, motivo...)."""
    g = df_glosas.copy()
    if "Operadora" not in g.columns:
        g["Operadora"] = g.attrs.get("operadora", "")
    ren = {"Data do evento": "Data do atendimento", "Nome do paciente": "Paciente",
           "Doc Original": "Guia do prestador"}
    g = g.rename(columns=ren)
    cols = ["Data do atendimento", "Paciente", "Guia do prestador", "Protocolo", "Lote", "Código",
            "Procedimento", "Quantidade", "Quantidade liberada", "Valor informado", "Valor pago",
            "Valor da glosa", "Status da glosa", "Motivo da glosa", "Data de pagamento", "Operadora"]
    cols = [c for c in cols if c in g.columns]
    g = g[cols].sort_values(["Operadora", "Data do atendimento"] if "Operadora" in cols else cols[:1]) \
        .reset_index(drop=True)

    total_glosa = float(g["Valor da glosa"].sum()) if len(g) else 0.0
    apresentado = float(df_completo["Valor informado"].sum()) if df_completo is not None and len(df_completo) \
        else float(g["Valor informado"].sum())
    guias = g["Guia do prestador"].replace("-", pd.NA).nunique() if "Guia do prestador" in g else 0
    datas = pd.to_datetime(g["Data do atendimento"], format="%d/%m/%Y", errors="coerce").dropna()
    periodo = f"{datas.min():%d/%m/%Y} a {datas.max():%d/%m/%Y}" if len(datas) else "-"

    def agrupa(col, titulo):
        if col not in g.columns or g.empty:
            return dict(titulo=titulo, df=pd.DataFrame())
        base = g
        if col == "Motivo da glosa":  # variações do mesmo motivo/código ficam juntas
            base = g.assign(**{col: me.agrupar_motivos(g[col]).replace("", "Sem motivo informado")})
        t = base.groupby(col, sort=False).agg(Itens=("Valor da glosa", "count"),
                                           **{"Valor da glosa": ("Valor da glosa", "sum")}).reset_index()
        t = t.sort_values("Valor da glosa", ascending=False)
        t["% do total glosado"] = (t["Valor da glosa"] / total_glosa) if total_glosa else 0.0
        return dict(titulo=titulo, df=t, moeda=["Valor da glosa"], inteiros=["Itens"], pct=["% do total glosado"])

    info = [("Operadora(s)", ", ".join(sorted({str(x) for x in g["Operadora"] if str(x)})) or "-"),
            ("Período dos atendimentos", periodo),
            ("Arquivo(s) analisado(s)", ", ".join(arquivos) if arquivos else "-"),
            ("Emitido em", _agora())]
    kpis = [("Itens glosados", len(g), "int"), ("Guias com glosa", guias, "int"),
            ("Valor glosado", total_glosa, "moeda"), ("Valor apresentado (demonstrativo)", apresentado, "moeda"),
            ("% glosado sobre o apresentado", (total_glosa / apresentado) if apresentado else 0.0, "pct")]
    # Lote só aparece no PDF quando o demonstrativo informa; o protocolo aparece sempre
    tem_lote = "Lote" in g.columns and g["Lote"].astype(str).str.strip().replace({"-": "", "nan": ""}).ne("").any()
    pdf_cols = ["Data do atendimento", "Paciente", "Guia do prestador", "Protocolo"] + (["Lote"] if tem_lote else []) \
        + ["Procedimento", "Valor informado", "Valor pago", "Valor da glosa", "Motivo da glosa"]
    detalhe = dict(nome="Glosas", titulo="Relação de glosas", df=g, moeda=["Valor informado", "Valor pago", "Valor da glosa"],
                   inteiros=["Quantidade", "Quantidade liberada"],
                   datas=["Data do atendimento", "Data de pagamento"], total=["Valor informado", "Valor pago", "Valor da glosa"],
                   status="Status da glosa", pdf_colunas=pdf_cols)
    return dict(
        titulo="Relatório de Glosas", subtitulo="Itens glosados pelas operadoras, por guia e lote",
        info=info, kpis=kpis,
        resumos=[agrupa("Motivo da glosa", "Glosas por motivo"), agrupa("Operadora", "Glosas por operadora/plano"),
                 agrupa("Lote" if tem_lote else "Protocolo", "Glosas por lote" if tem_lote else "Glosas por protocolo")],
        detalhes=[detalhe])


def spec_faturamento(res, arquivos_fat=None, arquivos_demo=None):
    """Relatório 3: faturamento x recebimento."""
    por_plano = mc.resumir_recebimento(res, "Plano")
    por_lote = mc.resumir_recebimento(res, "Lote")
    por_sit = mc.resumir_recebimento(res, "Situação")
    pagos = res[res["Mês de pagamento"] != "-"]
    por_mes = mc.resumir_recebimento(pagos, "Mês de pagamento") if not pagos.empty else pd.DataFrame()
    cobr = mc.lista_de_cobranca(res)

    fat = float(res["Valor faturado"].sum()) if len(res) else 0.0
    pago = float(res["Valor pago"].sum()) if len(res) else 0.0
    glosa = float(res["Valor da glosa"].sum()) if len(res) else 0.0
    aberto = float(res["Valor em aberto"].sum()) if len(res) else 0.0
    prazo = res["Dias para pagamento"].dropna()
    sem_retorno = int(res["Situação"].str.startswith("Não pago (sem retorno").sum()) if len(res) else 0

    fm = dict(moeda=["Valor faturado", "Valor pago", "Valor da glosa", "Valor em aberto"],
              inteiros=["Guias", "Prazo médio (dias)", "Dias para pagamento"], pct=["% pago"])
    info = [("Arquivo(s) de faturamento", ", ".join(arquivos_fat) if arquivos_fat else "-"),
            ("Demonstrativo(s) de pagamento", ", ".join(arquivos_demo) if arquivos_demo else "-"),
            ("Lotes analisados", int(res["Lote"].nunique()) if len(res) else 0),
            ("Emitido em", _agora())]
    kpis = [("Guias faturadas", len(res), "int"), ("Valor faturado", fat, "moeda"),
            ("Valor pago", pago, "moeda"), ("Valor glosado", glosa, "moeda"),
            ("Valor em aberto (cobrar)", aberto, "moeda"),
            ("% do faturado já pago", (pago / fat) if fat else 0.0, "pct"),
            ("Prazo médio de pagamento (dias)", float(prazo.mean()) if len(prazo) else 0.0, "int"),
            ("Guias sem retorno no demonstrativo", sem_retorno, "int")]
    cols_conf = ["Lote", "Guia", "Paciente", "Plano", "Data atendimento", "Valor faturado", "Valor pago",
                 "Valor da glosa", "Valor em aberto", "Situação", "Motivo da glosa", "Data de envio",
                 "Data de pagamento", "Mês de pagamento", "Dias para pagamento", "Base do prazo",
                 "Critério de vínculo"]
    pdf_cols = ["Lote", "Guia", "Paciente", "Plano", "Valor faturado", "Valor pago", "Valor da glosa",
                "Valor em aberto", "Situação", "Mês de pagamento", "Dias para pagamento", "Motivo da glosa"]
    det_base = dict(moeda=fm["moeda"], inteiros=["Dias para pagamento"],
                    datas=["Data atendimento", "Data de envio", "Data de pagamento"],
                    total=["Valor faturado", "Valor pago", "Valor da glosa", "Valor em aberto"],
                    status="Situação", pdf_colunas=pdf_cols)
    return dict(
        titulo="Relatório de Faturamento x Recebimento",
        subtitulo="Conferência do que foi faturado em cada lote contra o que o plano pagou",
        info=info, kpis=kpis,
        resumos=[dict(titulo="Por situação do pagamento", df=por_sit, **fm),
                 dict(titulo="Por plano de saúde", df=por_plano, **fm),
                 dict(titulo="Por mês de pagamento", df=por_mes, **fm),
                 dict(titulo="Por lote", df=por_lote, **fm)],
        detalhes=[dict(nome="A cobrar", titulo="Guias a cobrar ao plano (não pagas, glosadas ou pagas a menor)",
                       df=cobr[cols_conf] if len(cobr) else cobr, **det_base),
                  dict(nome="Conferência por guia", titulo="Conferência completa por guia",
                       df=res[cols_conf], **det_base)])
